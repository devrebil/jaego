"""재고 관리 프로그램 - 사내 서버 배포형 로컬 웹앱.

실행:  python app.py   ->  http://localhost:5000
"""
import os
import sys
import io
import json
import time
import re
import secrets
import shutil
import sqlite3
import zipfile
import holidays as holiday_lib
try:
    import qrcode
    import qrcode.image.svg
except ImportError:
    qrcode = None
from functools import wraps
from datetime import datetime, timedelta

from flask import (Flask, request, jsonify, session, send_file,
                   render_template, redirect, url_for, send_from_directory)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

import db
from db import (get_db, write_audit, now_str, get_setting, set_setting,
                DB_DIR, UPLOAD_DIR, DOMAINS, DOMAIN_LABELS, db_path)

BASE_DIR = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, "frozen", False) else __file__))
SESSION_SECRET_FILE = os.path.join(BASE_DIR, ".session_secret")


def load_session_secret():
    configured = os.environ.get("INVENTORY_SESSION_SECRET", "").strip()
    if configured:
        return configured
    try:
        with open(SESSION_SECRET_FILE, "r", encoding="utf-8") as secret_file:
            saved = secret_file.read().strip()
        if saved:
            return saved
    except FileNotFoundError:
        pass

    generated = secrets.token_hex(32)
    try:
        with open(SESSION_SECRET_FILE, "x", encoding="utf-8") as secret_file:
            secret_file.write(generated)
    except FileExistsError:
        with open(SESSION_SECRET_FILE, "r", encoding="utf-8") as secret_file:
            return secret_file.read().strip() or generated
    return generated


app = Flask(__name__)
app.secret_key = load_session_secret()


# ── 부팅: DB 초기화 + 기본 admin 계정 ─────────────────────────
def bootstrap():
    db.init_all()
    conn = get_db("users")
    n = conn.execute("SELECT COUNT(*) c FROM users").fetchone()["c"]
    if n == 0:
        conn.execute(
            "INSERT INTO users(username,password_hash,name,birthday,dept,role,created_at) VALUES(?,?,?,?,?,?,?)",
            ("admin", generate_password_hash("admin"), "관리자", "", "관리", "admin", now_str()))
        conn.commit()
    conn.close()
    if get_setting("company") is None:
        set_setting("company", "우리회사")
    if get_setting("home_period") is None:
        set_setting("home_period", "day")
    if get_setting("show_attendance_on_cal") is None:
        set_setting("show_attendance_on_cal", "1")
    if get_setting("attendance_visibility_fix_v1") is None:
        set_setting("show_attendance_on_cal", "1")
        set_setting("attendance_visibility_fix_v1", "1")
    if get_setting("tx_type_colors") is None:
        set_setting("tx_type_colors", "1")
    _migrate_sale_voucher_inventory()


# ── 인증 데코레이터 ──────────────────────────────────────────
def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if "user" not in session:
            return jsonify(ok=False, error="인증 필요"), 401
        return f(*a, **k)
    return w

def admin_required(f):
    @wraps(f)
    def w(*a, **k):
        if session.get("role") != "admin":
            return jsonify(ok=False, error="관리자 권한 필요"), 403
        return f(*a, **k)
    return w

def cur_user():
    return session.get("user")


def rows_to_list(rows):
    return [dict(r) for r in rows]


def _mentioned_users(text):
    names = set(re.findall(r"@([\w가-힣]+)", text or ""))
    if not names:
        return []
    conn = get_db("users")
    marks = ",".join("?" for _ in names)
    rows = conn.execute(f"SELECT username FROM users WHERE username IN ({marks})", tuple(names)).fetchall()
    conn.close()
    return [r["username"] for r in rows if r["username"] != cur_user()]


DUMMY_MARKERS = ("SIM-", "회귀테스트", "회귀검증", "DEBUG")
DUMMY_DIRECT_TABLES = {
    "contacts": [("contacts", ("name", "biz_no", "ceo", "phone", "email", "address", "memo"))],
    "items": [("items", ("sku", "name", "spec", "location", "distributor", "manufacturer"))],
    "transactions": [("transactions", ("contact", "sku", "item_name", "spec", "note", "channel", "order_no", "shipment_no", "project_ref", "author"))],
    "calendar": [
        ("schedules", ("title", "memo", "author")),
        ("daily_todos", ("user", "content")),
        ("week_plans", ("user", "content")),
    ],
    "memo": [("memos", ("author", "tag", "content"))],
    "notice": [("notices", ("author", "title", "content"))],
}


def _dummy_where(columns):
    parts, params = [], []
    for col in columns:
        for marker in DUMMY_MARKERS:
            parts.append(f"COALESCE({col}, '') LIKE ?")
            params.append(f"%{marker}%")
    return " OR ".join(parts), params


def _dummy_count(conn, table, columns):
    where, params = _dummy_where(columns)
    return conn.execute(f"SELECT COUNT(*) c FROM {table} WHERE {where}", params).fetchone()["c"]


def _dummy_delete(conn, table, columns):
    where, params = _dummy_where(columns)
    cur = conn.execute(f"DELETE FROM {table} WHERE {where}", params)
    return cur.rowcount if cur.rowcount >= 0 else 0


def _ids_for_dummy(conn, table, columns):
    where, params = _dummy_where(columns)
    return [r["id"] for r in conn.execute(f"SELECT id FROM {table} WHERE {where}", params).fetchall()]


def _in_clause(ids):
    return ",".join("?" for _ in ids)


def _delete_by_ids(conn, table, ids):
    ids = sorted(set(ids))
    if not ids:
        return 0
    cur = conn.execute(f"DELETE FROM {table} WHERE id IN ({_in_clause(ids)})", ids)
    return cur.rowcount if cur.rowcount >= 0 else 0


def _dummy_summary_counts():
    counts = {}
    for domain, tables in DUMMY_DIRECT_TABLES.items():
        conn = get_db(domain)
        counts[domain] = sum(_dummy_count(conn, table, cols) for table, cols in tables)
        conn.close()

    conn = get_db("projects")
    project_ids = _ids_for_dummy(conn, "projects", ("name", "status"))
    card_ids = _ids_for_dummy(conn, "cards", ("author", "content"))
    if project_ids:
        card_ids += [r["id"] for r in conn.execute(
            f"SELECT id FROM cards WHERE project_id IN ({_in_clause(project_ids)})",
            project_ids).fetchall()]
    card_ids = sorted(set(card_ids))
    related = len(project_ids) + len(card_ids)
    if card_ids:
        related += conn.execute(
            f"SELECT COUNT(*) c FROM comments WHERE card_id IN ({_in_clause(card_ids)})",
            card_ids).fetchone()["c"]
    related += _dummy_count(conn, "comments", ("author", "content"))
    if project_ids:
        related += conn.execute(
            f"SELECT COUNT(*) c FROM notifications WHERE project_id IN ({_in_clause(project_ids)})",
            project_ids).fetchone()["c"]
    related += _dummy_count(conn, "notifications", ("user", "actor", "message"))
    counts["projects"] = related
    conn.close()
    return counts


def _delete_project_dummy(conn):
    project_ids = _ids_for_dummy(conn, "projects", ("name", "status"))
    card_ids = _ids_for_dummy(conn, "cards", ("author", "content"))
    if project_ids:
        card_ids += [r["id"] for r in conn.execute(
            f"SELECT id FROM cards WHERE project_id IN ({_in_clause(project_ids)})",
            project_ids).fetchall()]
    card_ids = sorted(set(card_ids))

    comment_ids = _ids_for_dummy(conn, "comments", ("author", "content"))
    if card_ids:
        comment_ids += [r["id"] for r in conn.execute(
            f"SELECT id FROM comments WHERE card_id IN ({_in_clause(card_ids)})",
            card_ids).fetchall()]

    notification_ids = _ids_for_dummy(conn, "notifications", ("user", "actor", "message"))
    if project_ids:
        notification_ids += [r["id"] for r in conn.execute(
            f"SELECT id FROM notifications WHERE project_id IN ({_in_clause(project_ids)})",
            project_ids).fetchall()]

    return (
        _delete_by_ids(conn, "comments", comment_ids)
        + _delete_by_ids(conn, "notifications", notification_ids)
        + _delete_by_ids(conn, "cards", card_ids)
        + _delete_by_ids(conn, "projects", project_ids)
    )


def _create_mention_notifications(text, project_id, source_type, source_id):
    users = _mentioned_users(text)
    if not users:
        return
    conn = get_db("projects")
    for username in users:
        conn.execute("""INSERT INTO notifications(
            user,actor,project_id,source_type,source_id,message,dismissed,created_at
        ) VALUES(?,?,?,?,?,?,0,?)""",
                     (username, cur_user(), project_id, source_type, source_id,
                      (text or "")[:180], now_str()))
    conn.commit()
    conn.close()


# ── 페이지 ───────────────────────────────────────────────────
@app.route("/")
def index():
    return render_template("index.html")

@app.route("/uploads/<path:fn>")
def uploaded(fn):
    return send_from_directory(UPLOAD_DIR, fn)


def save_upload(file, prefix):
    if not file or not file.filename:
        return None
    stem, ext = os.path.splitext(file.filename)
    safe_stem = secure_filename(stem) or prefix
    safe_ext = re.sub(r"[^a-z0-9.]", "", ext.lower())[:10]
    fn = f"{prefix}_{int(time.time()*1000)}_{safe_stem}{safe_ext}"
    file.save(os.path.join(UPLOAD_DIR, fn))
    return fn


# ════════════════════════════════════════════════════════════
#  인증 / 계정
# ════════════════════════════════════════════════════════════
def mobile_access_url():
    """모바일에서 접속할 HTTPS 주소 (LAN IP 우선순위: 192.168 > 172 > 기타)."""
    try:
        from https_gateway import lan_ips
        ips = [ip for ip in lan_ips() if ip != "127.0.0.1"]
        ips.sort(key=lambda ip: 0 if ip.startswith("192.168.") else 1 if ip.startswith("172.") else 2)
        return f"https://{ips[0]}:5443" if ips else ""
    except Exception:
        return ""


@app.route("/api/me")
def api_me():
    if "user" not in session:
        return jsonify(logged_in=False, company=get_setting("company"))
    conn = get_db("users")
    user = conn.execute(
        "SELECT email,security_question FROM users WHERE username=?", (session["user"],)).fetchone()
    conn.close()
    return jsonify(logged_in=True, user=session["user"], name=session.get("name"),
                   role=session.get("role"), email=user["email"] if user else "",
                   security_question=user["security_question"] if user else "",
                   company=get_setting("company"),
                   home_period=get_setting("home_period"),
                   tx_type_colors=get_setting("tx_type_colors", "1"),
                   mobile_url=mobile_access_url())

@app.route("/api/mobile-qr")
@login_required
def api_mobile_qr():
    if qrcode is None:
        return jsonify(ok=False, error="QR 코드 기능을 사용할 수 없습니다."), 500
    url = mobile_access_url()
    if not url:
        return jsonify(ok=False, error="모바일 접속 주소를 확인할 수 없습니다."), 404
    img = qrcode.make(url, image_factory=qrcode.image.svg.SvgPathImage, box_size=10, border=2)
    buf = io.BytesIO()
    img.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="image/svg+xml")

@app.route("/api/users")
@login_required
def api_users():
    conn = get_db("users")
    rows = conn.execute("SELECT username,name,role,dept FROM users ORDER BY username").fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))


@app.route("/api/admin/users")
@admin_required
def api_admin_users():
    """관리자 회원관리용 목록. 비밀번호/복구답변 해시는 의도적으로 조회하지 않는다."""
    conn = get_db("users")
    rows = conn.execute(
        "SELECT username,name,email,dept,role,created_at FROM users ORDER BY name,username"
    ).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))

@app.route("/api/login", methods=["POST"])
def api_login():
    d = request.json or {}
    conn = get_db("users")
    u = conn.execute("SELECT * FROM users WHERE username=?", (d.get("username", "").strip(),)).fetchone()
    conn.close()
    if not u or not check_password_hash(u["password_hash"], d.get("password", "")):
        return jsonify(ok=False, error="아이디 또는 비밀번호가 올바르지 않습니다.")
    session["user"] = u["username"]; session["name"] = u["name"]; session["role"] = u["role"]
    ip = request.headers.get("X-Forwarded-For", request.remote_addr or "")
    write_audit(u["username"], "users", "login", u["username"], detail="로그인", ip=ip)
    return jsonify(ok=True, user=u["username"], name=u["name"], role=u["role"])

@app.route("/api/logout", methods=["POST"])
def api_logout():
    ip = request.headers.get("X-Forwarded-For", request.remote_addr or "")
    write_audit(cur_user(), "users", "logout", cur_user(), detail="로그아웃", ip=ip)
    session.clear()
    return jsonify(ok=True)

@app.route("/api/audit/client-event", methods=["POST"])
def api_audit_client_event():
    if "user" not in session:
        return jsonify(ok=False)
    d = request.json or {}
    ip = request.headers.get("X-Forwarded-For", request.remote_addr or "")
    write_audit(cur_user(), "users", d.get("action", "client"), None,
                detail=d.get("detail", ""), ip=ip)
    return jsonify(ok=True)

@app.route("/api/register", methods=["POST"])
def api_register():
    d = request.json or {}
    un = (d.get("username") or "").strip()
    email = (d.get("email") or "").strip().lower()
    question = (d.get("security_question") or "").strip()
    answer = (d.get("security_answer") or "").strip().lower()
    if not un or not d.get("password") or not email or not question or not answer:
        return jsonify(ok=False, error="아이디, 비밀번호, 이메일, 찾기 질문과 답변은 필수입니다.")
    if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
        return jsonify(ok=False, error="올바른 이메일 형식을 입력해주세요.")
    if len(d["password"]) < 6:
        return jsonify(ok=False, error="비밀번호는 6자 이상 입력해주세요.")
    conn = get_db("users")
    if conn.execute("SELECT 1 FROM users WHERE username=?", (un,)).fetchone():
        conn.close()
        return jsonify(ok=False, error="이미 존재하는 아이디입니다.")
    if conn.execute("SELECT 1 FROM users WHERE lower(email)=?", (email,)).fetchone():
        conn.close()
        return jsonify(ok=False, error="이미 등록된 이메일입니다.")
    try:
        conn.execute("""INSERT INTO users(
            username,password_hash,name,birthday,dept,role,created_at,
            email,security_question,security_answer_hash
        ) VALUES(?,?,?,?,?,?,?,?,?,?)""",
                     (un, generate_password_hash(d["password"]), d.get("name", ""), d.get("birthday", ""),
                      d.get("dept", ""), "member", now_str(), email, question,
                      generate_password_hash(answer)))
        conn.commit()
    except Exception:
        conn.close()
        return jsonify(ok=False, error="이미 존재하는 아이디입니다.")
    conn.close()
    write_audit(un, "users", "register", un, after={"name": d.get("name"), "dept": d.get("dept")}, detail="회원가입")
    return jsonify(ok=True)

@app.route("/api/account/update", methods=["POST"])
@login_required
def api_account_update():
    d = request.json or {}
    conn = get_db("users")
    u = conn.execute("SELECT * FROM users WHERE username=?", (cur_user(),)).fetchone()
    new_pw = d.get("password")
    new_un = (d.get("username") or u["username"]).strip()
    email = (d.get("email") if d.get("email") is not None else u["email"] or "").strip().lower()
    question = (d.get("security_question") if d.get("security_question") is not None
                else u["security_question"] or "").strip()
    answer = (d.get("security_answer") or "").strip().lower()
    pwh = generate_password_hash(new_pw) if new_pw else u["password_hash"]
    answer_hash = generate_password_hash(answer) if answer else u["security_answer_hash"]
    if email and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
        conn.close()
        return jsonify(ok=False, error="올바른 이메일 형식을 입력해주세요.")
    duplicate = conn.execute(
        "SELECT 1 FROM users WHERE lower(email)=? AND id<>?", (email, u["id"])).fetchone() if email else None
    if duplicate:
        conn.close()
        return jsonify(ok=False, error="이미 등록된 이메일입니다.")
    try:
        conn.execute("""UPDATE users SET username=?,password_hash=?,name=?,birthday=?,dept=?,
                     email=?,security_question=?,security_answer_hash=? WHERE id=?""",
                     (new_un, pwh, d.get("name", u["name"]), d.get("birthday", u["birthday"]),
                      d.get("dept", u["dept"]), email, question, answer_hash, u["id"]))
        conn.commit()
    except Exception:
        conn.close(); return jsonify(ok=False, error="아이디 중복")
    conn.close()
    session["user"] = new_un; session["name"] = d.get("name", u["name"])
    write_audit(new_un, "users", "update_account", u["id"], detail="계정 정보 수정")
    return jsonify(ok=True)


@app.route("/api/account/find-id", methods=["POST"])
def api_account_find_id():
    d = request.json or {}
    name = (d.get("name") or "").strip()
    email = (d.get("email") or "").strip().lower()
    conn = get_db("users")
    row = conn.execute(
        "SELECT username FROM users WHERE name=? AND lower(email)=?", (name, email)).fetchone()
    conn.close()
    if not row:
        return jsonify(ok=False, error="일치하는 계정을 찾을 수 없습니다.")
    return jsonify(ok=True, username=row["username"])


@app.route("/api/account/security-question", methods=["POST"])
def api_account_security_question():
    d = request.json or {}
    username = (d.get("username") or "").strip()
    email = (d.get("email") or "").strip().lower()
    conn = get_db("users")
    row = conn.execute(
        "SELECT security_question,security_answer_hash FROM users "
        "WHERE username=? AND lower(email)=?", (username, email)).fetchone()
    conn.close()
    if not row or not row["security_question"] or not row["security_answer_hash"]:
        return jsonify(ok=False, error="일치하는 계정 또는 등록된 찾기 질문이 없습니다.")
    return jsonify(ok=True, question=row["security_question"])


@app.route("/api/account/reset-password", methods=["POST"])
def api_account_reset_password():
    d = request.json or {}
    username = (d.get("username") or "").strip()
    email = (d.get("email") or "").strip().lower()
    answer = (d.get("security_answer") or "").strip().lower()
    new_password = d.get("new_password") or ""
    if len(new_password) < 6:
        return jsonify(ok=False, error="새 비밀번호는 6자 이상 입력해주세요.")
    conn = get_db("users")
    row = conn.execute(
        "SELECT id,security_answer_hash FROM users WHERE username=? AND lower(email)=?",
        (username, email)).fetchone()
    if not row or not row["security_answer_hash"] or not check_password_hash(
            row["security_answer_hash"], answer):
        conn.close()
        return jsonify(ok=False, error="보안 답변이 올바르지 않습니다.")
    conn.execute("UPDATE users SET password_hash=? WHERE id=?",
                 (generate_password_hash(new_password), row["id"]))
    conn.commit()
    conn.close()
    write_audit(username, "users", "reset_password", row["id"], detail="보안 질문으로 비밀번호 재설정")
    return jsonify(ok=True)

@app.route("/api/users/role", methods=["POST"])
@admin_required
def api_user_role():
    d = request.json or {}
    conn = get_db("users")
    conn.execute("UPDATE users SET role=? WHERE username=?", (d.get("role"), d.get("username")))
    conn.commit(); conn.close()
    write_audit(cur_user(), "users", "set_role", d.get("username"), after={"role": d.get("role")}, detail="권한 변경")
    return jsonify(ok=True)


# ════════════════════════════════════════════════════════════
#  전표 (판매입력/견적서, 기존 구매 데이터 호환)
# ════════════════════════════════════════════════════════════
VCH_KINDS = {"sale", "purchase", "quote"}
VCH_KIND_LABELS = {"sale": "판매", "purchase": "구매", "quote": "견적"}


def _vch_calc(vat_mode, qty, price):
    try:
        qty = float(qty or 0)
    except (TypeError, ValueError):
        qty = 0
    try:
        price = float(price or 0)
    except (TypeError, ValueError):
        price = 0
    amt = qty * price
    if vat_mode == "incl":      # 부가세 포함가 → 역산
        total = round(amt)
        supply = round(total / 1.1)
        return qty, price, supply, total - supply
    if vat_mode == "excl":      # 부가세 별도
        supply = round(amt)
        return qty, price, supply, round(amt * 0.1)
    return qty, price, round(amt), 0  # 미적용


def _delete_sale_voucher_transactions(vid, voucher=None, clean_legacy=True):
    """판매전표에 연결된 매출 행을 삭제하고 기존 수동 반영 행도 한 번 정리한다."""
    tx = get_db("transactions")
    linked = tx.execute(
        "SELECT COUNT(*) c FROM transactions WHERE source_voucher_id=?", (vid,)
    ).fetchone()["c"]
    deleted = tx.execute(
        "DELETE FROM transactions WHERE source_voucher_id=?", (vid,)
    ).rowcount
    if clean_legacy and not linked and voucher and voucher["posted"]:
        note = f"전표 {voucher['date']} -{voucher['seq']}"
        cur = tx.execute(
            """DELETE FROM transactions
               WHERE source_voucher_id IS NULL AND type='매출'
                 AND contact=? AND note=?""",
            (voucher["contact"], note),
        )
        deleted += max(0, cur.rowcount)
    tx.commit()
    tx.close()
    return deleted


def _sync_sale_voucher_transactions(vid, clean_legacy=True):
    """판매전표의 현재 품목을 재고 계산용 매출 행과 일치시킨다."""
    conn = get_db("vouchers")
    voucher = conn.execute("SELECT * FROM vouchers WHERE id=?", (vid,)).fetchone()
    if not voucher or voucher["kind"] != "sale":
        conn.close()
        return 0
    items = conn.execute(
        "SELECT * FROM voucher_items WHERE voucher_id=? ORDER BY sort,id", (vid,)
    ).fetchall()
    _delete_sale_voucher_transactions(vid, voucher, clean_legacy)
    tx = get_db("transactions")
    for item in items:
        tx.execute(
            """INSERT INTO transactions(
                date,type,contact,sku,item_name,spec,unit,
                buy_qty,buy_price,sell_qty,sell_price,return_qty,return_price,
                total,note,author,created_at,source_voucher_id,source_voucher_item_id)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                voucher["date"], "매출", voucher["contact"], item["sku"], item["name"],
                item["spec"], item["unit"], 0, 0, item["qty"], item["price"], 0, 0,
                item["supply"] + item["tax"], f"전표 {voucher['date']} -{voucher['seq']}",
                voucher["author"] or cur_user() or "system", now_str(), vid, item["id"],
            ),
        )
    tx.commit()
    tx.close()
    conn.execute(
        "UPDATE vouchers SET posted=1, updated_at=? WHERE id=?", (now_str(), vid)
    )
    conn.commit()
    conn.close()
    return len(items)


def _migrate_sale_voucher_inventory():
    """기존 판매전표도 자동 재고 반영 체계로 한 번 이관한다."""
    if get_setting("sale_voucher_auto_stock_v1") == "1":
        return
    conn = get_db("vouchers")
    ids = [r["id"] for r in conn.execute(
        "SELECT id FROM vouchers WHERE kind='sale' ORDER BY id"
    ).fetchall()]
    conn.close()
    for vid in ids:
        _sync_sale_voucher_transactions(vid, clean_legacy=True)
    set_setting("sale_voucher_auto_stock_v1", "1")


@app.route("/api/vouchers")
@login_required
def api_vouchers():
    kind = request.args.get("kind", "sale")
    conn = get_db("vouchers")
    rows = [dict(r) for r in conn.execute(
        "SELECT * FROM vouchers WHERE kind=? ORDER BY date DESC, seq DESC, id DESC", (kind,)).fetchall()]
    for v in rows:
        it = conn.execute("SELECT name, spec FROM voucher_items WHERE voucher_id=? ORDER BY sort, id",
                          (v["id"],)).fetchall()
        s = ""
        if it:
            s = it[0]["name"] + (f" [{it[0]['spec']}]" if it[0]["spec"] else "")
            if len(it) > 1:
                s += f" 외 {len(it) - 1}건"
        v["item_summary"] = s
    conn.close()
    return jsonify(rows)


@app.route("/api/vouchers/<int:vid>")
@login_required
def api_voucher_get(vid):
    conn = get_db("vouchers")
    v = conn.execute("SELECT * FROM vouchers WHERE id=?", (vid,)).fetchone()
    if not v:
        conn.close()
        return jsonify(ok=False, error="전표를 찾을 수 없습니다.")
    d = dict(v)
    d["ok"] = True
    d["items"] = [dict(r) for r in conn.execute(
        "SELECT * FROM voucher_items WHERE voucher_id=? ORDER BY sort, id", (vid,)).fetchall()]
    conn.close()
    return jsonify(d)


@app.route("/api/vouchers/save", methods=["POST"])
@login_required
def api_voucher_save():
    d = request.json or {}
    kind = d.get("kind")
    if kind not in VCH_KINDS:
        return jsonify(ok=False, error="잘못된 전표 종류입니다.")
    date = (d.get("date") or now_str()[:10]).strip()
    vat_mode = d.get("vat_mode") or "none"
    items = [x for x in (d.get("items") or [])
             if (x.get("name") or "").strip() or (x.get("sku") or "").strip()]
    if not items:
        return jsonify(ok=False, error="품목을 1건 이상 입력하세요.")
    if not (d.get("contact") or "").strip():
        return jsonify(ok=False, error="거래처를 입력하세요.")

    sup_t = tax_t = 0
    norm = []
    for i, x in enumerate(items):
        qty, price, sup, tax = _vch_calc(vat_mode, x.get("qty"), x.get("price"))
        sup_t += sup
        tax_t += tax
        norm.append(((x.get("sku") or "").strip(), (x.get("name") or "").strip(),
                     x.get("spec") or "", x.get("unit") or "",
                     qty, price, sup, tax, x.get("note") or "", i))

    conn = get_db("vouchers")
    vid = d.get("id")
    if vid:
        old = conn.execute("SELECT kind, date, seq FROM vouchers WHERE id=?", (vid,)).fetchone()
        if not old:
            conn.close()
            return jsonify(ok=False, error="전표를 찾을 수 없습니다.")
        if old["kind"] != kind:
            conn.close()
            return jsonify(ok=False, error="전표 종류는 변경할 수 없습니다.")
        seq = old["seq"]
        if old["date"] != date:  # 일자가 바뀌면 새 일자의 다음 번호 부여
            seq = (conn.execute("SELECT MAX(seq) m FROM vouchers WHERE kind=? AND date=?",
                                (kind, date)).fetchone()["m"] or 0) + 1
        conn.execute("""UPDATE vouchers SET date=?, seq=?, warehouse=?, contact=?, dept=?, vat_mode=?,
            memo=?, valid_until=?, supply=?, tax=?, total=?, updated_at=? WHERE id=?""",
            (date, seq, d.get("warehouse") or "", d.get("contact") or "", d.get("dept") or "",
             vat_mode, d.get("memo") or "", d.get("valid_until") or "",
             sup_t, tax_t, sup_t + tax_t, now_str(), vid))
        conn.execute("DELETE FROM voucher_items WHERE voucher_id=?", (vid,))
        action = "update"
    else:
        seq = (conn.execute("SELECT MAX(seq) m FROM vouchers WHERE kind=? AND date=?",
                            (kind, date)).fetchone()["m"] or 0) + 1
        cur = conn.execute("""INSERT INTO vouchers(kind, date, seq, warehouse, contact, dept, vat_mode,
            memo, valid_until, status, posted, printed, supply, tax, total, author, created_at, updated_at)
            VALUES(?,?,?,?,?,?,?,?,?,?,0,0,?,?,?,?,?,?)""",
            (kind, date, seq, d.get("warehouse") or "", d.get("contact") or "", d.get("dept") or "",
             vat_mode, d.get("memo") or "", d.get("valid_until") or "", d.get("status") or "미확인",
             sup_t, tax_t, sup_t + tax_t, cur_user(), now_str(), now_str()))
        vid = cur.lastrowid
        action = "create"
    for row in norm:
        conn.execute("""INSERT INTO voucher_items(voucher_id, sku, name, spec, unit, qty, price,
            supply, tax, note, sort) VALUES(?,?,?,?,?,?,?,?,?,?,?)""", (vid,) + row)
    conn.commit()
    conn.close()
    if kind == "sale":
        try:
            _sync_sale_voucher_transactions(vid, clean_legacy=True)
        except Exception as exc:
            return jsonify(ok=False, error=f"판매전표는 저장됐지만 재고 반영에 실패했습니다: {exc}"), 500
    label = VCH_KIND_LABELS[kind]
    write_audit(cur_user(), "vouchers", action, vid,
                detail=f"{label}전표 {date} -{seq} {d.get('contact') or ''} {len(norm)}품목 합계 {sup_t + tax_t:,.0f}")
    return jsonify(ok=True, id=vid, seq=seq)


@app.route("/api/vouchers/delete", methods=["POST"])
@login_required
def api_voucher_delete():
    vid = (request.json or {}).get("id")
    conn = get_db("vouchers")
    v = conn.execute("SELECT * FROM vouchers WHERE id=?", (vid,)).fetchone()
    if not v:
        conn.close()
        return jsonify(ok=False, error="전표를 찾을 수 없습니다.")
    if v["kind"] == "sale":
        try:
            _delete_sale_voucher_transactions(vid, v, clean_legacy=True)
        except Exception as exc:
            conn.close()
            return jsonify(ok=False, error=f"재고 복원에 실패해 판매전표를 삭제하지 않았습니다: {exc}"), 500
    conn.execute("DELETE FROM voucher_items WHERE voucher_id=?", (vid,))
    conn.execute("DELETE FROM vouchers WHERE id=?", (vid,))
    conn.commit()
    conn.close()
    write_audit(cur_user(), "vouchers", "delete", vid, before=dict(v),
                detail=f"{VCH_KIND_LABELS.get(v['kind'], '')}전표 삭제 {v['date']} -{v['seq']} {v['contact']}")
    return jsonify(ok=True)


@app.route("/api/vouchers/field", methods=["POST"])
@login_required
def api_voucher_field():
    d = request.json or {}
    field = d.get("field")
    if field not in ("status", "printed"):
        return jsonify(ok=False, error="허용되지 않은 필드")
    conn = get_db("vouchers")
    conn.execute(f"UPDATE vouchers SET {field}=?, updated_at=? WHERE id=?",
                 (d.get("value"), now_str(), d.get("id")))
    conn.commit()
    conn.close()
    if field == "status":
        write_audit(cur_user(), "vouchers", "status", d.get("id"), detail=f"상태 → {d.get('value')}")
    return jsonify(ok=True)


@app.route("/api/vouchers/post-tx", methods=["POST"])
@login_required
def api_voucher_post_tx():
    """판매/구매 전표를 매입·매출(회계) 장부로 반영한다."""
    vid = (request.json or {}).get("id")
    conn = get_db("vouchers")
    v = conn.execute("SELECT * FROM vouchers WHERE id=?", (vid,)).fetchone()
    if not v:
        conn.close()
        return jsonify(ok=False, error="전표를 찾을 수 없습니다.")
    if v["kind"] not in ("sale", "purchase"):
        conn.close()
        return jsonify(ok=False, error="판매/구매 전표만 회계 반영할 수 있습니다.")
    if v["kind"] == "sale":
        conn.close()
        count = _sync_sale_voucher_transactions(vid, clean_legacy=True)
        return jsonify(ok=True, count=count, automatic=True)
    if v["posted"]:
        conn.close()
        return jsonify(ok=False, error="이미 회계 반영된 전표입니다.")
    items = conn.execute("SELECT * FROM voucher_items WHERE voucher_id=? ORDER BY sort, id",
                         (vid,)).fetchall()
    tx = get_db("transactions")
    for it in items:
        if v["kind"] == "sale":
            row = (v["date"], "매출", v["contact"], it["sku"], it["name"], it["spec"], it["unit"],
                   0, 0, it["qty"], it["price"], 0, 0)
        else:
            row = (v["date"], "매입", v["contact"], it["sku"], it["name"], it["spec"], it["unit"],
                   it["qty"], it["price"], 0, 0, 0, 0)
        tx.execute("""INSERT INTO transactions(date, type, contact, sku, item_name, spec, unit,
            buy_qty, buy_price, sell_qty, sell_price, return_qty, return_price, total, note, author, created_at)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (*row, it["supply"] + it["tax"], f"전표 {v['date']} -{v['seq']}", cur_user(), now_str()))
    tx.commit()
    tx.close()
    conn.execute("UPDATE vouchers SET posted=1, updated_at=? WHERE id=?", (now_str(), vid))
    conn.commit()
    conn.close()
    write_audit(cur_user(), "vouchers", "post_tx", vid,
                detail=f"{VCH_KIND_LABELS[v['kind']]}전표 회계반영 {v['date']} -{v['seq']} {v['contact']} {len(items)}품목")
    return jsonify(ok=True, count=len(items))


# ════════════════════════════════════════════════════════════
#  설정 / 회사 / 백업·복원 / 히스토리
# ════════════════════════════════════════════════════════════
@app.route("/api/settings")
@login_required
def api_settings():
    keys = ["company", "home_period", "show_attendance_on_cal", "tx_type_colors",
            "company_biz_no", "company_ceo", "company_addr", "company_tel", "company_fax",
            "company_bank", "company_homepage", "print_footer_order", "print_color", "print_theme",
            "company_logo", "company_watermark", "company_watermark_size", "company_stamp", "company_stamp_position"]
    return jsonify({k: get_setting(k) for k in keys})

@app.route("/api/settings/save", methods=["POST"])
@login_required
def api_settings_save():
    d = request.json or {}
    # 회사명/권한 변경은 admin만
    admin_keys = {"company", "company_biz_no", "company_ceo", "company_addr", "company_tel", "company_fax",
                  "company_bank", "company_homepage", "print_footer_order", "print_color", "print_theme",
                  "company_watermark_size", "company_stamp_position"}
    for k, v in d.items():
        if k in admin_keys and session.get("role") != "admin":
            continue
        if k == "company_stamp_position" and v not in {"company", "ceo"}:
            continue
        if k == "company_watermark_size" and v not in {"small", "medium", "large"}:
            continue
        set_setting(k, v)
    write_audit(cur_user(), "users", "update_settings", detail=json.dumps(d, ensure_ascii=False))
    return jsonify(ok=True)


@app.route("/api/settings/print-assets", methods=["POST"])
@login_required
def api_settings_print_assets():
    if session.get("role") != "admin":
        return jsonify(ok=False, error="관리자만 출력 이미지를 변경할 수 있습니다."), 403
    saved = {}
    allowed_extensions = {".png", ".jpg", ".jpeg", ".webp"}
    for field, key, prefix in (("logo", "company_logo", "companylogo"),
                               ("watermark", "company_watermark", "companywatermark"),
                               ("stamp", "company_stamp", "companystamp")):
        file = request.files.get(field)
        if file and file.filename:
            ext = os.path.splitext(file.filename)[1].lower()
            if ext not in allowed_extensions:
                return jsonify(ok=False, error="PNG, JPG, JPEG, WEBP 이미지 파일만 업로드할 수 있습니다."), 400
            filename = save_upload(file, prefix)
            set_setting(key, filename)
            saved[key] = filename
    if not saved:
        return jsonify(ok=False, error="저장할 로고, 워터마크 또는 도장 이미지를 선택하세요."), 400
    write_audit(cur_user(), "users", "update_print_assets", detail="견적서/명세표 로고·워터마크·도장 이미지 변경")
    return jsonify(ok=True, **saved)


@app.route("/api/preferences")
@login_required
def api_preferences():
    """Per-user dashboard preferences shared by every PC/browser."""
    username = cur_user()
    result = {}
    for name, default in (("shortcuts", {}), ("quick_launch", []), ("search_sites", {}),
                          ("search_default_mode", "general"),
                          ("search_links", [{"name": "홈택스", "url": "https://www.hometax.go.kr/"}]),
                          ("week_start", 0), ("show_sidebar_shortcuts", True)):
        raw = get_setting(f"user:{username}:{name}")
        try:
            result[name] = json.loads(raw) if raw else default
        except (TypeError, ValueError):
            result[name] = default
    return jsonify(result)


@app.route("/api/preferences/save", methods=["POST"])
@login_required
def api_preferences_save():
    data = request.json or {}
    allowed = {"shortcuts", "quick_launch", "search_sites", "search_default_mode", "search_links",
               "week_start", "show_sidebar_shortcuts"}
    if "search_default_mode" in data:
        data["search_default_mode"] = (data.get("search_default_mode")
                                       if data.get("search_default_mode") in {"general", "all"}
                                       else "general")
    if "search_links" in data:
        if not isinstance(data["search_links"], list):
            return jsonify(ok=False, error="멀티검색 바로가기 형식이 올바르지 않습니다."), 400
        if len(data["search_links"]) > 5:
            return jsonify(ok=False, error="멀티검색 바로가기는 최대 5개까지 저장할 수 있습니다."), 400
        links = []
        for item in data["search_links"]:
            if not isinstance(item, dict):
                continue
            name = str(item.get("name") or "").strip()[:40]
            url = str(item.get("url") or "").strip()[:500]
            if not name and not url:
                continue
            if not name or not re.match(r"^https?://", url, re.I):
                return jsonify(ok=False, error="바로가기 이름과 http:// 또는 https:// 주소를 확인하세요."), 400
            links.append({"name": name, "url": url})
        data["search_links"] = links
    saved = {}
    for name in allowed:
        if name in data:
            value = data[name]
            set_setting(f"user:{cur_user()}:{name}", json.dumps(value, ensure_ascii=False))
            saved[name] = value
    write_audit(cur_user(), "users", "update_preferences", detail="사용자 화면 설정 변경")
    return jsonify(ok=True, saved=saved)


@app.route("/api/leave-days")
@login_required
def api_leave_days():
    conn = get_db("users")
    users = conn.execute("SELECT username,name FROM users ORDER BY name,username").fetchall()
    conn.close()
    if session.get("role") != "admin":
        users = [u for u in users if u["username"] == cur_user()]
    return jsonify([{"username": u["username"], "name": u["name"],
                     "days": float(get_setting(f"leave_days:{u['username']}", "0") or 0)}
                    for u in users])


@app.route("/api/leave-days/save", methods=["POST"])
@admin_required
def api_leave_days_save():
    data = request.json or {}
    username = (data.get("username") or "").strip()
    try:
        days = max(0, float(data.get("days") or 0))
    except (TypeError, ValueError):
        return jsonify(ok=False, error="연차 일수는 숫자로 입력하세요."), 400
    conn = get_db("users")
    exists = conn.execute("SELECT 1 FROM users WHERE username=?", (username,)).fetchone()
    conn.close()
    if not exists:
        return jsonify(ok=False, error="사용자를 찾을 수 없습니다."), 404
    set_setting(f"leave_days:{username}", days)
    write_audit(cur_user(), "users", "set_leave_days", username,
                after={"days": days}, detail="사용자 연차 일수 변경")
    return jsonify(ok=True)

@app.route("/api/backup/<domain>")
@admin_required
def api_backup(domain):
    if domain not in DOMAINS:
        return jsonify(ok=False, error="알 수 없는 DB"), 400
    write_audit(cur_user(), "audit", "backup", domain, detail=f"{domain} 백업")
    return send_file(db_path(domain), as_attachment=True, download_name=DOMAINS[domain])

@app.route("/api/restore/<domain>", methods=["POST"])
@admin_required
def api_restore(domain):
    if domain not in DOMAINS:
        return jsonify(ok=False, error="알 수 없는 DB"), 400
    f = request.files.get("file")
    if not f:
        return jsonify(ok=False, error="파일이 없습니다.")
    # 안전: 기존 파일 .bak 백업 후 덮어쓰기
    p = db_path(domain)
    if os.path.exists(p):
        os.replace(p, p + ".bak")
    f.save(p)
    write_audit(cur_user(), "audit", "restore", domain, detail=f"{domain} 복원")
    return jsonify(ok=True)

@app.route("/api/backup/list")
@admin_required
def api_backup_list():
    out = []
    for dom, fn in DOMAINS.items():
        p = db_path(dom)
        out.append({"domain": dom, "label": DOMAIN_LABELS[dom], "file": fn,
                    "size": os.path.getsize(p) if os.path.exists(p) else 0})
    return jsonify(out)

@app.route("/api/system/backup-all", methods=["POST"])
def api_backup_all():
    # 이 PC(로컬 프로세스)에서만 호출 가능. 종료 스크립트가 로그인 세션 없이
    # 프로그램 종료 직전 전체 DB를 백업하기 위한 내부 전용 엔드포인트.
    if request.remote_addr not in ("127.0.0.1", "::1"):
        return jsonify(ok=False, error="로컬에서만 호출할 수 있습니다."), 403
    dest_dir, copied = db.backup_all(tag=(request.args.get("tag") or "manual"))
    return jsonify(ok=True, dir=os.path.basename(dest_dir), files=copied)

_AUTO_BACKUP_TAG_LABELS = {"auto": "자동(10분)", "exit": "종료 시", "shutdown": "Windows 종료 시", "manual": "수동", "before_restore": "복원 전 스냅샷"}

def _auto_backup_entry(folder_name):
    path = os.path.join(db.AUTO_BACKUP_DIR, folder_name)
    # 폴더명 형식: YYYYMMDD_HHMMSSffffff_tag (마이크로초까지 포함, 옛 백업은 초 단위만 있을 수 있음)
    parts = folder_name.split("_", 2)
    tag = parts[2] if len(parts) == 3 else "manual"
    time_label = folder_name
    if len(parts) == 3:
        for fmt in ("%Y%m%d%H%M%S%f", "%Y%m%d%H%M%S"):
            try:
                dt = datetime.strptime(parts[0] + parts[1], fmt)
                time_label = dt.strftime("%Y-%m-%d %H:%M:%S")
                break
            except ValueError:
                continue
    files = [f for f in os.listdir(path) if os.path.isfile(os.path.join(path, f))]
    size = sum(os.path.getsize(os.path.join(path, f)) for f in files)
    return {"name": folder_name, "time": time_label, "tag": tag,
            "tag_label": _AUTO_BACKUP_TAG_LABELS.get(tag, tag),
            "file_count": len(files), "size": size}

@app.route("/api/backup/auto/list")
@admin_required
def api_backup_auto_list():
    if not os.path.isdir(db.AUTO_BACKUP_DIR):
        return jsonify([])
    names = sorted(
        (n for n in os.listdir(db.AUTO_BACKUP_DIR) if os.path.isdir(os.path.join(db.AUTO_BACKUP_DIR, n))),
        reverse=True,
    )
    return jsonify([_auto_backup_entry(n) for n in names])

@app.route("/api/backup/auto/run", methods=["POST"])
@admin_required
def api_backup_auto_run():
    dest_dir, copied = db.backup_all(tag="manual")
    write_audit(cur_user(), "audit", "backup_all", os.path.basename(dest_dir),
                detail=f"전체 DB 수동 백업 ({len(copied)}개)")
    return jsonify(ok=True, entry=_auto_backup_entry(os.path.basename(dest_dir)))

@app.route("/api/backup/auto/<name>/download")
@admin_required
def api_backup_auto_download(name):
    safe_name = os.path.basename(name)
    path = os.path.join(db.AUTO_BACKUP_DIR, safe_name)
    if safe_name != name or not os.path.isdir(path):
        return jsonify(ok=False, error="백업을 찾을 수 없습니다."), 404
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for fn in os.listdir(path):
            fp = os.path.join(path, fn)
            if os.path.isfile(fp):
                z.write(fp, fn)
    buf.seek(0)
    write_audit(cur_user(), "audit", "backup_download", safe_name, detail="자동 백업 다운로드")
    return send_file(buf, as_attachment=True, download_name=f"{safe_name}.zip", mimetype="application/zip")

@app.route("/api/backup/auto/<name>/restore", methods=["POST"])
@admin_required
def api_backup_auto_restore(name):
    safe_name = os.path.basename(name)
    path = os.path.join(db.AUTO_BACKUP_DIR, safe_name)
    if safe_name != name or not os.path.isdir(path):
        return jsonify(ok=False, error="백업을 찾을 수 없습니다."), 404

    # 안전: 복원으로 덮어쓰기 전에 현재 상태를 먼저 스냅샷으로 남긴다.
    db.backup_all(tag="before_restore")

    filename_to_domain = {fn: dom for dom, fn in DOMAINS.items()}
    restored = []
    for fn in os.listdir(path):
        domain = filename_to_domain.get(fn)
        if not domain:
            continue
        shutil.copy2(os.path.join(path, fn), db_path(domain))
        restored.append(domain)

    write_audit(cur_user(), "audit", "backup_restore_all", safe_name,
                detail=f"자동 백업 전체 복원 ({len(restored)}개 DB)")
    return jsonify(ok=True, restored=restored)

@app.route("/api/dummy-data/summary")
@admin_required
def api_dummy_data_summary():
    counts = _dummy_summary_counts()
    rows = [{"domain": dom, "label": DOMAIN_LABELS.get(dom, dom), "count": counts.get(dom, 0)}
            for dom in sorted(counts)]
    return jsonify(ok=True, markers=list(DUMMY_MARKERS), total=sum(counts.values()), rows=rows)

@app.route("/api/dummy-data/delete", methods=["POST"])
@admin_required
def api_dummy_data_delete():
    data = request.json or {}
    if data.get("confirm") != "DELETE_DUMMY":
        return jsonify(ok=False, error="확인 문구가 일치하지 않습니다."), 400

    before = _dummy_summary_counts()
    if sum(before.values()) == 0:
        return jsonify(ok=True, total=0, rows=[], backups=[])

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    touched_domains = [dom for dom, count in before.items() if count > 0]
    backups = []
    for dom in touched_domains:
        src = db_path(dom)
        if os.path.exists(src):
            dst = f"{src}.dummy_cleanup_{stamp}.bak"
            shutil.copy2(src, dst)
            backups.append(os.path.basename(dst))

    deleted = {}
    for domain, tables in DUMMY_DIRECT_TABLES.items():
        if before.get(domain, 0) <= 0:
            continue
        conn = get_db(domain)
        try:
            deleted[domain] = sum(_dummy_delete(conn, table, cols) for table, cols in tables)
            conn.commit()
        finally:
            conn.close()

    if before.get("projects", 0) > 0:
        conn = get_db("projects")
        try:
            deleted["projects"] = _delete_project_dummy(conn)
            conn.commit()
        finally:
            conn.close()

    rows = [{"domain": dom, "label": DOMAIN_LABELS.get(dom, dom), "count": deleted.get(dom, 0)}
            for dom in sorted(deleted)]
    total = sum(deleted.values())
    write_audit(cur_user(), "audit", "delete_dummy_data",
                after={"deleted": deleted, "backups": backups},
                detail=f"더미 데이터 삭제 {total}건")
    return jsonify(ok=True, total=total, rows=rows, backups=backups)


# 계정과 users.db의 설정(회사/사업자/출력/사용자 환경설정)은 보존하고
# 실제 업무 기록만 비우는 완전 초기화 대상이다. 태그/거래유형도 설정 성격이라 유지한다.
BUSINESS_RESET_TABLES = {
    "contacts": ["contacts"],
    "items": ["items"],
    "transactions": ["transactions"],
    "calendar": ["daily_todos", "holidays", "schedules", "week_plans"],
    "projects": ["comments", "notifications", "cards", "projects"],
    "memo": ["memos"],
    "attendance": ["attendance"],
    "notice": ["notices"],
    "vouchers": ["voucher_items", "vouchers"],
    "audit": ["audit"],
}


def reset_business_data(make_backup=True):
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = os.path.join(db.DB_DIR, f"full_reset_backup_{stamp}")

    # 회사 로고/워터마크/도장은 사업자 설정 자산이므로 완전 초기화에서도 보존한다.
    settings_conn = get_db("users")
    try:
        asset_rows = settings_conn.execute(
            "SELECT value FROM settings WHERE key IN "
            "('company_logo','company_watermark','company_stamp')"
        ).fetchall()
        preserved_assets = {
            os.path.basename(row["value"])
            for row in asset_rows
            if row["value"]
        }
    finally:
        settings_conn.close()

    if make_backup:
        os.makedirs(backup_dir, exist_ok=False)
        for domain in BUSINESS_RESET_TABLES:
            src = db_path(domain)
            if os.path.exists(src):
                shutil.copy2(src, os.path.join(backup_dir, os.path.basename(src)))
        if os.path.isdir(db.UPLOAD_DIR):
            upload_backup_dir = os.path.join(backup_dir, "uploads")
            os.makedirs(upload_backup_dir, exist_ok=True)
            for filename in os.listdir(db.UPLOAD_DIR):
                src = os.path.join(db.UPLOAD_DIR, filename)
                if os.path.isfile(src):
                    shutil.copy2(src, os.path.join(upload_backup_dir, filename))

    deleted = {}
    for domain, tables in BUSINESS_RESET_TABLES.items():
        conn = get_db(domain)
        domain_count = 0
        try:
            for table in tables:
                count = conn.execute(f"SELECT COUNT(*) c FROM {table}").fetchone()["c"]
                conn.execute(f"DELETE FROM {table}")
                domain_count += count
                try:
                    conn.execute("DELETE FROM sqlite_sequence WHERE name=?", (table,))
                except sqlite3.OperationalError:
                    pass
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
        deleted[domain] = domain_count

    deleted_uploads = 0
    if os.path.isdir(db.UPLOAD_DIR):
        for filename in os.listdir(db.UPLOAD_DIR):
            path = os.path.join(db.UPLOAD_DIR, filename)
            if os.path.isfile(path) and filename not in preserved_assets:
                os.remove(path)
                deleted_uploads += 1
    return {
        "deleted": deleted,
        "total": sum(deleted.values()),
        "deleted_uploads": deleted_uploads,
        "backup_dir": os.path.basename(backup_dir) if make_backup else None,
    }


@app.route("/api/business-data/reset", methods=["POST"])
@admin_required
def api_business_data_reset():
    data = request.json or {}
    if data.get("confirm") != "RESET_ALL_DATA":
        return jsonify(ok=False, error="확인 문구가 일치하지 않습니다."), 400
    try:
        result = reset_business_data(make_backup=True)
    except Exception:
        return jsonify(ok=False, error="초기화 중 오류가 발생했습니다. 데이터는 백업되어 있습니다."), 500
    return jsonify(ok=True, **result)

@app.route("/api/audit")
@admin_required
def api_audit():
    domain = request.args.get("domain", "")
    user = request.args.get("user", "")
    start = request.args.get("start", "")
    end = request.args.get("end", "")
    q = "SELECT * FROM audit WHERE 1=1"
    p = []
    if domain: q += " AND domain=?"; p.append(domain)
    if user: q += " AND user=?"; p.append(user)
    if start: q += " AND ts>=?"; p.append(start)
    if end: q += " AND ts<=?"; p.append(end + " 23:59:59")
    q += " ORDER BY id DESC LIMIT 500"
    conn = get_db("audit")
    rows = conn.execute(q, p).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))


# ════════════════════════════════════════════════════════════
#  출퇴근
# ════════════════════════════════════════════════════════════
@app.route("/api/attendance", methods=["POST"])
@login_required
def api_attendance():
    d = request.json or {}
    t = d.get("type")  # 'in' | 'out'
    if t not in ("in", "out"):
        return jsonify(ok=False, error="출퇴근 구분을 확인하세요."), 400
    ts = now_str()
    conn = get_db("attendance")
    today = ts[:10]
    rows = conn.execute(
        "SELECT * FROM attendance WHERE user=? AND ts>=? AND ts<=? ORDER BY ts",
        (cur_user(), today, today + " 23:59:59")).fetchall()
    has_in = any(r["type"] == "in" for r in rows)
    has_out = any(r["type"] == "out" for r in rows)
    if t == "in" and has_in:
        conn.close()
        return jsonify(ok=False, error="오늘은 이미 출근 처리되었습니다."), 400
    if t == "out" and not has_in:
        conn.close()
        return jsonify(ok=False, error="출근 기록이 있어야 퇴근 처리할 수 있습니다."), 400
    if t == "out" and has_out:
        conn.close()
        return jsonify(ok=False, error="오늘은 이미 퇴근 처리되었습니다."), 400
    conn.execute("INSERT INTO attendance(user,type,ts) VALUES(?,?,?)", (cur_user(), t, ts))
    conn.commit(); conn.close()
    write_audit(cur_user(), "attendance", "checkin" if t == "in" else "checkout",
                detail=f"{'출근' if t=='in' else '퇴근'} {ts}")
    return jsonify(ok=True, ts=ts, status=_attendance_today_status(cur_user()))

def _attendance_today_status(username):
    today = datetime.now().strftime("%Y-%m-%d")
    conn = get_db("attendance")
    rows = conn.execute(
        "SELECT * FROM attendance WHERE user=? AND ts>=? AND ts<=? ORDER BY ts",
        (username, today, today + " 23:59:59")).fetchall()
    conn.close()
    check_in = next((r["ts"] for r in rows if r["type"] == "in"), "")
    check_out = next((r["ts"] for r in reversed(rows) if r["type"] == "out"), "")
    return {
        "date": today,
        "check_in": check_in,
        "check_out": check_out,
        "can_check_in": not check_in,
        "can_check_out": bool(check_in) and not check_out,
    }

@app.route("/api/attendance/today")
@login_required
def api_attendance_today():
    return jsonify(_attendance_today_status(cur_user()))

@app.route("/api/attendance/list")
@login_required
def api_attendance_list():
    conn = get_db("attendance")
    if session.get("role") == "admin":
        rows = conn.execute("SELECT * FROM attendance ORDER BY id DESC LIMIT 200").fetchall()
    else:
        rows = conn.execute("SELECT * FROM attendance WHERE user=? ORDER BY id DESC LIMIT 200", (cur_user(),)).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))


@app.route("/api/attendance/history")
@admin_required
def api_attendance_history():
    user = request.args.get("user", "").strip()
    conn = get_db("attendance")
    if user:
        rows = conn.execute(
            "SELECT * FROM attendance WHERE user=? ORDER BY ts DESC LIMIT 2000", (user,)).fetchall()
    else:
        rows = conn.execute("SELECT * FROM attendance ORDER BY ts DESC LIMIT 2000").fetchall()
    conn.close()

    grouped = {}
    for row in rows:
        item = dict(row)
        day = (item.get("ts") or "")[:10]
        key = (item.get("user"), day)
        group = grouped.setdefault(key, {
            "user": item.get("user"), "date": day, "check_in": "", "check_out": "", "records": 0
        })
        group["records"] += 1
        ts = item.get("ts") or ""
        if item.get("type") == "in" and (not group["check_in"] or ts < group["check_in"]):
            group["check_in"] = ts
        if item.get("type") == "out" and (not group["check_out"] or ts > group["check_out"]):
            group["check_out"] = ts

    result = []
    for group in grouped.values():
        minutes = None
        if group["check_in"] and group["check_out"]:
            try:
                start = datetime.strptime(group["check_in"][:19], "%Y-%m-%d %H:%M:%S")
                end = datetime.strptime(group["check_out"][:19], "%Y-%m-%d %H:%M:%S")
                minutes = max(0, int((end - start).total_seconds() // 60))
            except ValueError:
                pass
        group["work_minutes"] = minutes
        result.append(group)
    result.sort(key=lambda x: (x["date"], x["user"]), reverse=True)
    return jsonify(result)


# ════════════════════════════════════════════════════════════
#  거래처
# ════════════════════════════════════════════════════════════
@app.route("/api/contacts")
@login_required
def api_contacts():
    q = request.args.get("q", "")
    conn = get_db("contacts")
    rows = conn.execute("SELECT * FROM contacts WHERE name LIKE ? OR biz_no LIKE ? ORDER BY name",
                        (f"%{q}%", f"%{q}%")).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))

@app.route("/api/contacts/save", methods=["POST"])
@login_required
def api_contacts_save():
    d = request.form
    photo = save_upload(request.files.get("photo"), "contact")
    biz_cert = save_upload(request.files.get("biz_cert"), "bizcert")
    conn = get_db("contacts")
    cid = d.get("id")
    if cid:
        old = dict(conn.execute("SELECT * FROM contacts WHERE id=?", (cid,)).fetchone())
        conn.execute("""UPDATE contacts SET name=?,biz_no=?,ceo=?,phone=?,email=?,address=?,memo=?,
            photo=COALESCE(?,photo), biz_cert=COALESCE(?,biz_cert) WHERE id=?""",
            (d.get("name"), d.get("biz_no"), d.get("ceo"), d.get("phone"), d.get("email"),
             d.get("address"), d.get("memo"), photo, biz_cert, cid))
        conn.commit()
        write_audit(cur_user(), "contacts", "update", cid, before=old, detail="거래처 수정")
    else:
        cur = conn.execute("""INSERT INTO contacts(name,biz_no,ceo,phone,email,address,memo,photo,biz_cert,created_at)
            VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (d.get("name"), d.get("biz_no"), d.get("ceo"), d.get("phone"), d.get("email"),
             d.get("address"), d.get("memo"), photo, biz_cert, now_str()))
        conn.commit(); cid = cur.lastrowid
        write_audit(cur_user(), "contacts", "create", cid, after={"name": d.get("name")}, detail="거래처 등록")
    conn.close()
    return jsonify(ok=True, id=cid)

@app.route("/api/contacts/delete", methods=["POST"])
@login_required
def api_contacts_delete():
    cid = (request.json or {}).get("id")
    conn = get_db("contacts")
    old = conn.execute("SELECT * FROM contacts WHERE id=?", (cid,)).fetchone()
    conn.execute("DELETE FROM contacts WHERE id=?", (cid,))
    conn.commit(); conn.close()
    write_audit(cur_user(), "contacts", "delete", cid, before=dict(old) if old else None, detail="거래처 삭제")
    return jsonify(ok=True)


# ════════════════════════════════════════════════════════════
#  품목
# ════════════════════════════════════════════════════════════
@app.route("/api/items")
@login_required
def api_items():
    q = request.args.get("q", "")
    conn = get_db("items")
    rows = conn.execute("SELECT * FROM items WHERE sku LIKE ? OR name LIKE ? ORDER BY sku",
                        (f"%{q}%", f"%{q}%")).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))

@app.route("/api/items/save", methods=["POST"])
@login_required
def api_items_save():
    d = request.form
    photo = save_upload(request.files.get("photo"), "item")
    conn = get_db("items")
    iid = d.get("id")
    try:
        if iid:
            old = dict(conn.execute("SELECT * FROM items WHERE id=?", (iid,)).fetchone())
            conn.execute("""UPDATE items SET sku=?,name=?,spec=?,unit=?,location=?,buy_price=?,sell_price=?,
                distributor=?,manufacturer=?,photo=COALESCE(?,photo) WHERE id=?""",
                (d.get("sku"), d.get("name"), d.get("spec"), d.get("unit"), d.get("location"),
                 float(d.get("buy_price") or 0), float(d.get("sell_price") or 0),
                 d.get("distributor", ""), d.get("manufacturer", ""), photo, iid))
            conn.commit()
            write_audit(cur_user(), "items", "update", iid, before=old, detail="품목 수정")
        else:
            cur = conn.execute("""INSERT INTO items(sku,name,spec,unit,location,buy_price,sell_price,
                distributor,manufacturer,photo,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (d.get("sku"), d.get("name"), d.get("spec"), d.get("unit"), d.get("location"),
                 float(d.get("buy_price") or 0), float(d.get("sell_price") or 0),
                 d.get("distributor", ""), d.get("manufacturer", ""), photo, now_str()))
            conn.commit(); iid = cur.lastrowid
            write_audit(cur_user(), "items", "create", iid, after={"sku": d.get("sku"), "name": d.get("name")}, detail="품목 등록")
    except Exception as e:
        conn.close(); return jsonify(ok=False, error=f"저장 실패(품목코드 중복?): {e}")
    conn.close()
    return jsonify(ok=True, id=iid)

@app.route("/api/items/delete", methods=["POST"])
@login_required
def api_items_delete():
    iid = (request.json or {}).get("id")
    conn = get_db("items")
    old = conn.execute("SELECT * FROM items WHERE id=?", (iid,)).fetchone()
    conn.execute("DELETE FROM items WHERE id=?", (iid,))
    conn.commit(); conn.close()
    write_audit(cur_user(), "items", "delete", iid, before=dict(old) if old else None, detail="품목 삭제")
    return jsonify(ok=True)

@app.route("/api/items/import", methods=["POST"])
@login_required
def api_items_import():
    """엑셀(xlsx)/csv 일괄 업로드: 품목코드|품목명|규격|단위|위치|매입단가|판매단가|유통사|제조사"""
    f = request.files.get("file")
    if not f:
        return jsonify(ok=False, error="파일이 없습니다.")
    rows = []
    name = f.filename.lower()
    try:
        if name.endswith(".csv"):
            import csv
            txt = f.read().decode("utf-8-sig")
            for r in csv.reader(io.StringIO(txt)):
                rows.append(r)
        else:
            import openpyxl
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            ws = wb.active
            for r in ws.iter_rows(values_only=True):
                rows.append(list(r))
    except Exception as e:
        return jsonify(ok=False, error=f"파일 읽기 실패: {e}")

    conn = get_db("items")
    cnt = 0
    for i, r in enumerate(rows):
        if not r or all(c is None or str(c).strip() == "" for c in r):
            continue
        # 헤더로 보이면 스킵
        head = str(r[0]).strip().lower()
        if i == 0 and ("sku" in head or "상품코드" in head or "코드" in head):
            continue
        raw = list(r)
        if len(raw) <= 6:
            sku, nm, spec, unit, bp, sp = (raw + [None] * 6)[:6]
            location = distributor = manufacturer = ""
        else:
            sku, nm, spec, unit, location, bp, sp, distributor, manufacturer = (raw + [None] * 9)[:9]
        if not sku:
            continue
        try:
            old = conn.execute("SELECT * FROM items WHERE sku=?", (str(sku).strip(),)).fetchone()
            conn.execute("""INSERT INTO items(sku,name,spec,unit,location,buy_price,sell_price,
                distributor,manufacturer,created_at)
                VALUES(?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(sku) DO UPDATE SET name=excluded.name,spec=excluded.spec,unit=excluded.unit,
                location=excluded.location,buy_price=excluded.buy_price,sell_price=excluded.sell_price,
                distributor=excluded.distributor,manufacturer=excluded.manufacturer""",
                (str(sku).strip(), str(nm or "").strip(), str(spec or "").strip(), str(unit or "").strip(),
                 str(location or "").strip(),
                 float(bp or 0), float(sp or 0),
                 str(distributor or "").strip(), str(manufacturer or "").strip(), now_str()))
            cnt += 1
            write_audit(cur_user(), "items", "import_update" if old else "import_create",
                        str(sku).strip(), before=dict(old) if old else None,
                        after={"sku": str(sku).strip(), "name": str(nm or "").strip()},
                        detail="품목 엑셀 업로드")
        except Exception:
            continue
    conn.commit(); conn.close()
    write_audit(cur_user(), "items", "import", detail=f"엑셀 일괄 등록 {cnt}건")
    return jsonify(ok=True, count=cnt)


# ════════════════════════════════════════════════════════════
#  매입·매출 전표
# ════════════════════════════════════════════════════════════
@app.route("/api/tx/types")
@login_required
def api_tx_types():
    conn = get_db("transactions")
    rows = conn.execute("SELECT name FROM tx_types ORDER BY id").fetchall()
    conn.close()
    return jsonify([r["name"] for r in rows])

@app.route("/api/tx/types/add", methods=["POST"])
@login_required
def api_tx_types_add():
    nm = (request.json or {}).get("name", "").strip()
    if nm:
        conn = get_db("transactions")
        conn.execute("INSERT OR IGNORE INTO tx_types(name) VALUES(?)", (nm,))
        conn.commit(); conn.close()
        write_audit(cur_user(), "transactions", "add_type", detail=f"구분 추가: {nm}")
    return jsonify(ok=True)


@app.route("/api/tx/types/delete", methods=["POST"])
@login_required
def api_tx_types_delete():
    nm = (request.json or {}).get("name", "").strip()
    protected = {"매입", "매출", "반품", "실사재고변경"}
    if not nm or nm in protected:
        return jsonify(ok=False, error="기본 구분은 삭제할 수 없습니다.")
    conn = get_db("transactions")
    used = conn.execute("SELECT COUNT(*) c FROM transactions WHERE type=?", (nm,)).fetchone()["c"]
    if used:
        conn.close()
        return jsonify(ok=False, error=f"이미 전표 {used}건에서 사용 중인 구분입니다.")
    conn.execute("DELETE FROM tx_types WHERE name=?", (nm,))
    conn.commit()
    conn.close()
    write_audit(cur_user(), "transactions", "delete_type", detail=f"구분 삭제: {nm}")
    return jsonify(ok=True)

@app.route("/api/tx")
@login_required
def api_tx():
    start = request.args.get("start", "")
    end = request.args.get("end", "")
    q = "SELECT * FROM transactions WHERE 1=1"; p = []
    if start: q += " AND date>=?"; p.append(start)
    if end: q += " AND date<=?"; p.append(end)
    contact = request.args.get("contact", "")
    if contact: q += " AND contact=?"; p.append(contact)
    tx_type = request.args.get("type", "")
    if tx_type: q += " AND type=?"; p.append(tx_type)
    # limit 기본 1000건(기존 동작), 0이면 전부, offset으로 이어서 조회. 전체 건수는 X-Total-Count 헤더로 전달
    limit = request.args.get("limit", default=1000, type=int)
    offset = max(0, request.args.get("offset", default=0, type=int))
    conn = get_db("transactions")
    total_count = conn.execute(q.replace("SELECT *", "SELECT COUNT(*)", 1), p).fetchone()[0]
    q += " ORDER BY date DESC, id DESC"
    if limit > 0:
        q += " LIMIT ? OFFSET ?"
        p = p + [limit, offset]
    rows = list(conn.execute(q, p).fetchall())
    focus_id = request.args.get("focus_id", type=int)
    if focus_id and not any(r["id"] == focus_id for r in rows):
        focus = conn.execute("SELECT * FROM transactions WHERE id=?", (focus_id,)).fetchone()
        if focus:
            rows.append(focus)
    conn.close()
    resp = jsonify(rows_to_list(rows))
    resp.headers["X-Total-Count"] = str(total_count)
    return resp

TX_FIELDS = ("date", "type", "contact", "sku", "item_name", "spec", "unit",
             "buy_qty", "buy_price", "sell_qty", "sell_price",
             "return_qty", "return_price", "note",
             "channel", "order_no", "shipment_no", "project_ref")

def _calc_total(d):
    return (float(d.get("buy_qty") or 0) * float(d.get("buy_price") or 0)
            + float(d.get("sell_qty") or 0) * float(d.get("sell_price") or 0)
            + float(d.get("return_qty") or 0) * float(d.get("return_price") or 0))


def _xlsx_response(title, headers, rows, filename):
    import openpyxl
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill, Alignment

    def clean(v):
        # 바코드 스캔 품목코드 등에 섞인 제어문자(GS1 구분자 \x1d 등)는 엑셀에 쓸 수 없어 제거
        return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4A90E2")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([clean(c) for c in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        width = min(40, max(10, max(len(str(c.value or "")) for c in col) + 2))
        ws.column_dimensions[col[0].column_letter].width = width
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/api/template/<kind>")
@login_required
def api_template(kind):
    if kind == "items":
        return _xlsx_response("품목 양식",
                              ["품목코드", "품목명", "규격", "단위", "위치", "매입단가", "판매단가", "유통사", "제조사"],
                              [["A-0001", "예시 품목", "규격", "개", "A-01", 0, 0, "예시유통", "예시제조"]],
                              "품목_등록_양식.xlsx")
    if kind == "tx":
        headers = ["일자", "구분", "거래처명", "품목코드", "품목명", "규격", "단위",
                   "매입수량", "매입단가", "매출수량", "매출단가",
                   "반품수량", "반품단가", "비고",
                   "채널", "주문번호", "송장번호", "프로젝트"]
        row = [datetime.now().strftime("%Y-%m-%d"), "매입", "", "A-0001", "예시 품목",
               "규격", "개", 1, 0, 0, 0, 0, 0, "", "", "", "", ""]
        return _xlsx_response("매입매출 양식", headers, [row], "매입매출_등록_양식.xlsx")
    return jsonify(ok=False, error="지원하지 않는 양식입니다."), 404


@app.route("/api/export/<kind>")
@login_required
def api_export(kind):
    if kind == "contacts":
        conn = get_db("contacts")
        rows = conn.execute("SELECT name,biz_no,ceo,phone,email,address,memo,created_at FROM contacts ORDER BY name").fetchall()
        conn.close()
        return _xlsx_response("거래처", ["거래처명", "사업자번호", "대표자", "연락처", "이메일", "주소", "메모", "등록시각"],
                              rows, "거래처.xlsx")
    if kind == "items":
        conn = get_db("items")
        rows = conn.execute("""SELECT sku,name,spec,unit,location,buy_price,sell_price,
            distributor,manufacturer,created_at FROM items ORDER BY sku""").fetchall()
        conn.close()
        return _xlsx_response("품목", ["품목코드", "품목명", "규격", "단위", "위치", "매입단가", "판매단가",
                              "유통사", "제조사", "등록시각"], rows, "품목.xlsx")
    if kind in ("tx", "inout"):
        conn = get_db("transactions")
        rows = conn.execute("""SELECT date,created_at,type,contact,sku,item_name,spec,unit,
            buy_qty,buy_price,sell_qty,sell_price,return_qty,return_price,total,note,
            channel,order_no,shipment_no,project_ref,author
            FROM transactions ORDER BY date DESC,id DESC""").fetchall()
        conn.close()
        # 품목의 유통사/제조사를 품목코드로 연결해 컬럼 추가
        it = get_db("items")
        info = {r["sku"]: (r["distributor"] or "", r["manufacturer"] or "")
                for r in it.execute("SELECT sku,distributor,manufacturer FROM items").fetchall()}
        it.close()
        rows = [list(r) + list(info.get(r["sku"], ("", ""))) for r in rows]
        return _xlsx_response("매입매출", ["일자", "정확한 기록시각", "구분", "거래처", "품목코드", "품목명", "규격", "단위",
                              "매입수량", "매입단가", "매출수량", "매출단가", "반품수량", "반품단가",
                              "합계", "비고", "채널", "주문번호", "송장번호", "프로젝트",
                              "작성자", "유통사", "제조사"], rows, "매입매출_입출고.xlsx")
    if kind == "inventory":
        data = _inventory()
        rows = [[x[k] for k in ("sku", "item_name", "spec", "unit", "location", "distributor", "manufacturer",
                                "in_qty", "out_qty", "ret_qty", "stock", "buy_amt", "sell_amt",
                                "gross_profit", "cash_flow")] for x in data]
        return _xlsx_response("재고현황", ["품목코드", "품목명", "규격", "단위", "위치", "유통사", "제조사",
                              "입고", "출고", "반품", "현재고", "매입액", "매출액",
                              "추정 매출총이익", "현금흐름"], rows, "재고현황.xlsx")
    if kind in ("month", "year"):
        mode = kind
        start_month = request.args.get("start_month", "")
        end_month = request.args.get("end_month", "")
        if start_month and end_month:
            sy, sm = map(int, start_month.split("-"))
            ey, em = map(int, end_month.split("-"))
            start = f"{sy:04d}-{sm:02d}-01"
            next_end = datetime(ey + (em == 12), 1 if em == 12 else em + 1, 1)
            end = (next_end - timedelta(days=1)).strftime("%Y-%m-%d")
            filename = f"{start_month}_{end_month}_입출현황.xlsx"
        else:
            year = int(request.args.get("year") or datetime.now().year)
            month = int(request.args.get("month") or datetime.now().month)
            filename = f"{year}_{month if mode == 'month' else '연간'}_입출현황.xlsx"
        if not start_month and mode == "month":
            start = f"{year:04d}-{month:02d}-01"
            next_month = datetime(year + (month == 12), 1 if month == 12 else month + 1, 1)
            end = (next_month - timedelta(days=1)).strftime("%Y-%m-%d")
        elif not start_month:
            start, end = f"{year:04d}-01-01", f"{year:04d}-12-31"
        data = _inventory(start, end)
        rows = [[x[k] for k in ("sku", "item_name", "buy_amt", "sell_amt", "gross_profit", "cash_flow")] for x in data]
        return _xlsx_response("입출현황", ["품목코드", "품목명", "매입액", "매출액", "추정 매출총이익", "현금흐름"],
                              rows, filename)
    return jsonify(ok=False, error="지원하지 않는 다운로드입니다."), 404


@app.route("/api/tx/import", methods=["POST"])
@login_required
def api_tx_import():
    f = request.files.get("file")
    if not f:
        return jsonify(ok=False, error="파일이 없습니다.")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
    except Exception as e:
        return jsonify(ok=False, error=f"파일 읽기 실패: {e}")
    conn = get_db("transactions")
    cnt = 0
    fields = TX_FIELDS
    for i, row in enumerate(rows):
        if i == 0:
            continue
        vals = (list(row) + [""] * len(fields))[:len(fields)]
        if not vals[0] or not vals[3]:
            continue
        if hasattr(vals[0], "strftime"):
            vals[0] = vals[0].strftime("%Y-%m-%d")
        d = dict(zip(fields, vals))
        total = _calc_total(d)
        cur = conn.execute(f"""INSERT INTO transactions({','.join(fields)},total,author,created_at)
            VALUES({','.join('?' for _ in fields)},?,?,?)""",
            (*vals, total, cur_user(), now_str()))
        cnt += 1
        write_audit(cur_user(), "transactions", "import_create", cur.lastrowid,
                    after={"type": d.get("type"), "sku": d.get("sku"), "total": total},
                    detail="매입매출 엑셀 업로드")
    conn.commit()
    conn.close()
    return jsonify(ok=True, count=cnt)

@app.route("/api/tx/save", methods=["POST"])
@login_required
def api_tx_save():
    d = request.json or {}
    total = _calc_total(d)
    conn = get_db("transactions")
    tid = d.get("id")
    fields = TX_FIELDS
    if tid:
        old_row = conn.execute("SELECT * FROM transactions WHERE id=?", (tid,)).fetchone()
        if not old_row:
            conn.close()
            return jsonify(ok=False, error="매입 내역을 찾을 수 없습니다."), 404
        if old_row["source_voucher_id"]:
            conn.close()
            return jsonify(ok=False, error="판매입력에서 자동 등록된 매출은 판매전표에서 수정하세요."), 400
        old = dict(old_row)
        conn.execute(f"""UPDATE transactions SET {','.join(f+'=?' for f in fields)}, total=? WHERE id=?""",
                     (*[d.get(f) for f in fields], total, tid))
        conn.commit()
        write_audit(cur_user(), "transactions", "update", tid, before=old, detail="전표 수정")
    else:
        cur = conn.execute(f"""INSERT INTO transactions({','.join(fields)},total,author,created_at)
            VALUES({','.join('?' for _ in fields)},?,?,?)""",
            (*[d.get(f) for f in fields], total, cur_user(), now_str()))
        conn.commit(); tid = cur.lastrowid
        write_audit(cur_user(), "transactions", "create", tid,
                    after={"type": d.get("type"), "sku": d.get("sku"), "total": total}, detail="전표 등록")
    conn.close()
    return jsonify(ok=True, id=tid, total=total)


@app.route("/api/tx/batch", methods=["POST"])
@login_required
def api_tx_batch():
    rows = (request.json or {}).get("items") or []
    if not isinstance(rows, list) or not rows:
        return jsonify(ok=False, error="처리할 품목이 없습니다."), 400
    if len(rows) > 500:
        return jsonify(ok=False, error="한 번에 최대 500개 품목까지 처리할 수 있습니다."), 400
    fields = TX_FIELDS
    conn = get_db("transactions")
    saved = []
    try:
        for data in rows:
            if not (data.get("sku") or "").strip() or not (data.get("item_name") or "").strip():
                raise ValueError("품목코드와 품목명을 확인하세요.")
            total = _calc_total(data)
            cur = conn.execute(f"""INSERT INTO transactions({','.join(fields)},total,author,created_at)
                VALUES({','.join('?' for _ in fields)},?,?,?)""",
                (*[data.get(f) for f in fields], total, cur_user(), now_str()))
            saved.append((cur.lastrowid, data, total))
        conn.commit()
    except Exception as exc:
        conn.rollback()
        conn.close()
        return jsonify(ok=False, error=f"일괄 처리 실패: {exc}"), 400
    conn.close()
    for transaction_id, data, total in saved:
        write_audit(cur_user(), "transactions", "create", transaction_id,
                    after={"type": data.get("type"), "sku": data.get("sku"), "total": total},
                    detail="바코드 일괄 입출고")
    return jsonify(ok=True, count=len(saved), ids=[row[0] for row in saved])

@app.route("/api/tx/delete", methods=["POST"])
@login_required
def api_tx_delete():
    tid = (request.json or {}).get("id")
    conn = get_db("transactions")
    old = conn.execute("SELECT * FROM transactions WHERE id=?", (tid,)).fetchone()
    if old and old["source_voucher_id"]:
        conn.close()
        return jsonify(ok=False, error="판매입력에서 자동 등록된 매출은 판매전표에서 삭제하세요."), 400
    conn.execute("DELETE FROM transactions WHERE id=?", (tid,))
    conn.commit(); conn.close()
    write_audit(cur_user(), "transactions", "delete", tid, before=dict(old) if old else None, detail="전표 삭제")
    return jsonify(ok=True)


@app.route("/api/inventory/adjust", methods=["POST"])
@login_required
def api_inventory_adjust():
    d = request.json or {}
    sku = (d.get("sku") or "").strip()
    desired = float(d.get("stock") or 0)
    current_row = next((x for x in _inventory() if x["sku"] == sku), None)
    current = float(current_row["stock"] if current_row else 0)
    delta = desired - current
    if not sku or delta == 0:
        return jsonify(ok=False, error="품목코드를 확인하거나 변경 수량을 입력하세요.")
    items = get_db("items")
    item = items.execute("SELECT * FROM items WHERE sku=?", (sku,)).fetchone()
    items.close()
    if not item:
        return jsonify(ok=False, error="등록된 품목이 아닙니다.")
    reason = (d.get("reason") or "").strip() or "실사재고변경"
    payload = {
        "date": datetime.now().strftime("%Y-%m-%d"),
        "type": "실사재고변경", "contact": "", "sku": sku,
        "item_name": item["name"], "spec": item["spec"], "unit": item["unit"],
        "buy_qty": delta if delta > 0 else 0, "buy_price": 0,
        "sell_qty": -delta if delta < 0 else 0, "sell_price": 0,
        "return_qty": 0, "return_price": 0,
        "note": reason, "channel": "", "order_no": "", "shipment_no": "", "project_ref": "",
    }
    conn = get_db("transactions")
    fields = TX_FIELDS
    cur = conn.execute(f"""INSERT INTO transactions({','.join(fields)},total,author,created_at)
        VALUES({','.join('?' for _ in fields)},?,?,?)""",
        (*[payload[f] for f in fields], 0, cur_user(), now_str()))
    conn.commit()
    tid = cur.lastrowid
    conn.close()
    write_audit(cur_user(), "inventory", "adjust", sku,
                before={"stock": current}, after={"stock": desired, "transaction_id": tid},
                detail=reason)
    return jsonify(ok=True, id=tid, before=current, after=desired)


# ════════════════════════════════════════════════════════════
#  재고 현황 / 월간·년간 입출 현황
# ════════════════════════════════════════════════════════════
def _inventory(start=None, end=None):
    q = "SELECT * FROM transactions WHERE 1=1"; p = []
    if start: q += " AND date>=?"; p.append(start)
    if end: q += " AND date<=?"; p.append(end)
    conn = get_db("transactions")
    rows = conn.execute(q, p).fetchall()
    conn.close()
    agg = {}
    items = get_db("items")
    for item in items.execute("SELECT sku,name,spec,unit,location,distributor,manufacturer,buy_price FROM items ORDER BY sku").fetchall():
        agg[item["sku"]] = {"sku": item["sku"], "item_name": item["name"], "spec": item["spec"],
                            "unit": item["unit"], "location": item["location"],
                            "distributor": item["distributor"] or "", "manufacturer": item["manufacturer"] or "",
                            "unit_cost": item["buy_price"] or 0,
                            "in_qty": 0, "out_qty": 0, "ret_qty": 0, "buy_amt": 0, "sell_amt": 0}
    items.close()
    for r in rows:
        sku = r["sku"] or "-"
        a = agg.setdefault(sku, {"sku": sku, "item_name": r["item_name"], "spec": r["spec"], "unit": r["unit"], "location": "",
                                 "distributor": "", "manufacturer": "", "unit_cost": 0,
                                 "in_qty": 0, "out_qty": 0, "ret_qty": 0, "buy_amt": 0, "sell_amt": 0})
        a["item_name"] = r["item_name"] or a["item_name"]
        a["in_qty"] += r["buy_qty"] or 0
        a["out_qty"] += r["sell_qty"] or 0
        a["ret_qty"] += r["return_qty"] or 0
        a["buy_amt"] += (r["buy_qty"] or 0) * (r["buy_price"] or 0)
        a["sell_amt"] += (r["sell_qty"] or 0) * (r["sell_price"] or 0)
    for a in agg.values():
        a["stock"] = a["in_qty"] - a["out_qty"] - a["ret_qty"]
        a["estimated_cogs"] = a["out_qty"] * (a.get("unit_cost") or 0)
        a["gross_profit"] = a["sell_amt"] - a["estimated_cogs"]
        a["cash_flow"] = a["sell_amt"] - a["buy_amt"]
        a["profit"] = a["gross_profit"]
    return list(agg.values())

def _report_summary(table):
    buy = sum(t["buy_amt"] for t in table)
    sell = sum(t["sell_amt"] for t in table)
    gross_profit = sum(t.get("gross_profit", t.get("profit", 0)) for t in table)
    cash_flow = sell - buy
    return {"buy": buy, "sell": sell, "profit": gross_profit,
            "gross_profit": gross_profit, "cash_flow": cash_flow}

@app.route("/api/inventory")
@login_required
def api_inventory():
    return jsonify(_inventory())

@app.route("/api/report")
@login_required
def api_report():
    """월간/년간 및 시작월~종료월 범위 입출 현황."""
    mode = request.args.get("mode", "month")
    start_month = request.args.get("start_month", "")
    end_month = request.args.get("end_month", "")
    if start_month and end_month:
        try:
            sy, sm = map(int, start_month.split("-"))
            ey, em = map(int, end_month.split("-"))
            start = f"{sy:04d}-{sm:02d}-01"
            next_end = datetime(ey + (em == 12), 1 if em == 12 else em + 1, 1)
            end = (next_end - timedelta(days=1)).strftime("%Y-%m-%d")
        except ValueError:
            return jsonify(ok=False, error="조회 기간이 올바르지 않습니다."), 400
        if start > end:
            return jsonify(ok=False, error="시작월은 종료월보다 빠르거나 같아야 합니다."), 400
        series = {}
        conn = get_db("transactions")
        rows = conn.execute("SELECT * FROM transactions WHERE date>=? AND date<=?", (start, end)).fetchall()
        conn.close()
        cursor = datetime(sy, sm, 1)
        finish = datetime(ey, em, 1)
        while cursor <= finish:
            key = cursor.strftime("%Y-%m")
            series[key] = {"label": key, "buy": 0, "sell": 0}
            cursor = datetime(cursor.year + (cursor.month == 12), 1 if cursor.month == 12 else cursor.month + 1, 1)
        for r in rows:
            key = (r["date"] or "")[:7]
            s = series.setdefault(key, {"label": key, "buy": 0, "sell": 0})
            s["buy"] += (r["buy_qty"] or 0) * (r["buy_price"] or 0)
            s["sell"] += (r["sell_qty"] or 0) * (r["sell_price"] or 0)
        table = _inventory(start, end)
        series_list = [series[k] for k in sorted(series)]
        return jsonify(table=table, series=series_list, range=True, start=start, end=end,
                       summary=_report_summary(table))
    year = int(request.args.get("year") or datetime.now().year)
    if mode == "month":
        month = int(request.args.get("month") or datetime.now().month)
        start = f"{year:04d}-{month:02d}-01"
        nm = month + 1; ny = year
        if nm > 12: nm = 1; ny += 1
        end = (datetime(ny, nm, 1) - timedelta(days=1)).strftime("%Y-%m-%d")
        # 일별 시리즈
        series = {}
        conn = get_db("transactions")
        rows = conn.execute("SELECT * FROM transactions WHERE date>=? AND date<=?", (start, end)).fetchall()
        conn.close()
        for r in rows:
            day = r["date"]
            s = series.setdefault(day, {"label": day, "buy": 0, "sell": 0})
            s["buy"] += (r["buy_qty"] or 0) * (r["buy_price"] or 0)
            s["sell"] += (r["sell_qty"] or 0) * (r["sell_price"] or 0)
    else:
        start = f"{year:04d}-01-01"; end = f"{year:04d}-12-31"
        series = {
            f"{month:02d}": {"label": f"{month}월", "buy": 0, "sell": 0}
            for month in range(1, 13)
        }
        conn = get_db("transactions")
        rows = conn.execute("SELECT * FROM transactions WHERE date>=? AND date<=?", (start, end)).fetchall()
        conn.close()
        for r in rows:
            m = (r["date"] or "")[5:7]
            s = series.setdefault(m, {"label": f"{int(m)}월" if m else "-", "buy": 0, "sell": 0})
            s["buy"] += (r["buy_qty"] or 0) * (r["buy_price"] or 0)
            s["sell"] += (r["sell_qty"] or 0) * (r["sell_price"] or 0)
    table = _inventory(start, end)
    series_list = [series[key] for key in sorted(series)]
    return jsonify(table=table, series=series_list,
                   summary=_report_summary(table))


# ════════════════════════════════════════════════════════════
#  캘린더 / 일정
# ════════════════════════════════════════════════════════════
@app.route("/api/schedules")
@login_required
def api_schedules():
    conn = get_db("calendar")
    rows = conn.execute("SELECT * FROM schedules ORDER BY start_date").fetchall()
    custom_holidays = conn.execute("SELECT * FROM holidays").fetchall()
    conn.close()
    current_year = datetime.now().year
    holiday_map = {
        day.isoformat(): name
        for day, name in holiday_lib.country_holidays(
            "KR", years=range(current_year - 10, current_year + 11), language="ko"
        ).items()
    }
    for holiday in custom_holidays:
        holiday_map[holiday["date"]] = holiday["name"]
    holiday_rows = [
        {"date": day, "name": name}
        for day, name in sorted(holiday_map.items())
    ]
    # 프로젝트 시작/완료를 일정 이벤트로 합침
    pj = get_db("projects")
    pjrows = pj.execute("SELECT * FROM projects").fetchall()
    pj.close()
    proj_events = []
    for p in pjrows:
        if p["start_date"]:
            proj_events.append({"title": f"📁 {p['name']} 시작", "start_date": p["start_date"],
                                "end_date": p["start_date"], "type": "project"})
        if p["end_date"]:
            proj_events.append({"title": f"✅ {p['name']} 완료", "start_date": p["end_date"],
                                "end_date": p["end_date"], "type": "project"})
    # 입출고(전표) 이벤트
    tx = get_db("transactions")
    txrows = tx.execute("SELECT date, type, item_name, COUNT(*) c FROM transactions GROUP BY date,type").fetchall()
    tx.close()
    tx_events = [{"title": f"{r['type']} {r['c']}건", "start_date": r["date"], "end_date": r["date"], "type": "tx"}
                 for r in txrows if r["date"]]
    # 출퇴근(관리자 토글 ON일 때만)
    att_events = []
    if session.get("role") == "admin" and get_setting("show_attendance_on_cal") == "1":
        at = get_db("attendance")
        atrows = at.execute("SELECT * FROM attendance ORDER BY id DESC LIMIT 300").fetchall()
        at.close()
        for r in atrows:
            d = (r["ts"] or "")[:10]
            att_events.append({"title": f"🕒 {r['user']} {'출근' if r['type']=='in' else '퇴근'} {(r['ts'] or '')[11:19]}",
                               "start_date": d, "end_date": d, "type": "attendance"})
    return jsonify(schedules=rows_to_list(rows), holidays=holiday_rows,
                   project_events=proj_events, tx_events=tx_events, attendance_events=att_events)

@app.route("/api/schedules/save", methods=["POST"])
@login_required
def api_schedules_save():
    d = request.json or {}
    conn = get_db("calendar")
    sid = d.get("id")
    if sid:
        old = dict(conn.execute("SELECT * FROM schedules WHERE id=?", (sid,)).fetchone())
        conn.execute("UPDATE schedules SET title=?,start_date=?,end_date=?,start_time=?,end_time=?,memo=?,status=?,color=? WHERE id=?",
                     (d.get("title"), d.get("start_date"), d.get("end_date") or d.get("start_date"),
                      d.get("start_time", ""), d.get("end_time", ""), d.get("memo"), d.get("status", "open"),
                      d.get("color") or "#2d7dd2", sid))
        conn.commit()
        write_audit(cur_user(), "calendar", "update", sid, before=old, detail="일정 수정")
    else:
        cur = conn.execute("INSERT INTO schedules(title,start_date,end_date,start_time,end_time,memo,status,author,created_at,color) VALUES(?,?,?,?,?,?,?,?,?,?)",
                           (d.get("title"), d.get("start_date"), d.get("end_date") or d.get("start_date"),
                            d.get("start_time", ""), d.get("end_time", ""), d.get("memo"), "open", cur_user(), now_str(),
                            d.get("color") or "#2d7dd2"))
        conn.commit(); sid = cur.lastrowid
        write_audit(cur_user(), "calendar", "create", sid, after={"title": d.get("title")}, detail="일정 등록")
    conn.close()
    return jsonify(ok=True, id=sid)

@app.route("/api/schedules/status", methods=["POST"])
@login_required
def api_schedules_status():
    d = request.json or {}
    conn = get_db("calendar")
    conn.execute("UPDATE schedules SET status=? WHERE id=?", (d.get("status"), d.get("id")))
    conn.commit(); conn.close()
    write_audit(cur_user(), "calendar", "status", d.get("id"), after={"status": d.get("status")}, detail="일정 상태 변경")
    return jsonify(ok=True)


@app.route("/api/calendar/todos")
@login_required
def api_calendar_todos():
    date = request.args.get("date", datetime.now().strftime("%Y-%m-%d"))
    conn = get_db("calendar")
    rows = conn.execute("SELECT * FROM daily_todos WHERE user=? AND date=? ORDER BY id",
                        (cur_user(), date)).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))


@app.route("/api/calendar/todos/save", methods=["POST"])
@login_required
def api_calendar_todos_save():
    d = request.json or {}
    conn = get_db("calendar")
    schedule_created = False
    saved_todo_date = d.get("date") or datetime.now().strftime("%Y-%m-%d")
    if d.get("id"):
        conn.execute("UPDATE daily_todos SET content=?,done=? WHERE id=? AND user=?",
                     (d.get("content", ""), 1 if d.get("done") else 0, d.get("id"), cur_user()))
        tid = d.get("id")
    else:
        content = (d.get("content") or "").strip()
        schedule = _schedule_from_todo(content, saved_todo_date)
        if schedule:
            saved_todo_date = schedule["start_date"]
        cur = conn.execute("INSERT INTO daily_todos(user,date,content,done,created_at) VALUES(?,?,?,?,?)",
                           (cur_user(), saved_todo_date, content, 0, now_str()))
        tid = cur.lastrowid
        if schedule:
            conn.execute(
                "INSERT INTO schedules(title,start_date,end_date,start_time,end_time,memo,status,author,created_at,color) "
                "VALUES(?,?,?,?,?,?,?,?,?,?)",
                (schedule["title"], schedule["start_date"], schedule["end_date"], schedule["start_time"],
                 schedule["end_time"], "투두리스트에서 자동 등록", "open", cur_user(), now_str(), "#2d7dd2"))
            schedule_created = True
    conn.commit()
    conn.close()
    return jsonify(ok=True, id=tid, todo_date=saved_todo_date, schedule_created=schedule_created)


TODO_KOREAN_HOURS = {
    "한": 1, "두": 2, "세": 3, "네": 4, "다섯": 5, "여섯": 6,
    "일곱": 7, "여덟": 8, "아홉": 9, "열": 10, "열한": 11, "열두": 12,
}
TODO_HOUR_PATTERN = r"(?:[01]?\d|2[0-3]|한|두|세|네|다섯|여섯|일곱|여덟|아홉|열두|열한|열)"


def _todo_clock_match(text):
    """문장 앞의 숫자·한글 시각을 분 단위로 해석한다."""
    colon = re.match(
        r"^\s*(?:(오전|오후)\s*)?([01]?\d|2[0-3]):([0-5]\d)(?=\s|$|부터|까지|에(?=\s|$)|[~～-])",
        text)
    if colon:
        period, hour_text, minute_text = colon.groups()
    else:
        korean = re.match(
            rf"^\s*(?:(오전|오후)\s*)?({TODO_HOUR_PATTERN})\s*시(?:\s*([0-5]?\d)\s*분)?(?=\s|$|부터|까지|에(?=\s|$)|[~～-])",
            text)
        if not korean:
            return None
        period, hour_text, minute_text = korean.groups()

    hour = TODO_KOREAN_HOURS.get(hour_text, int(hour_text) if hour_text.isdigit() else -1)
    minute = int(minute_text or 0)
    if period:
        if not 1 <= hour <= 12:
            return None
        hour = hour % 12 + (12 if period == "오후" else 0)
    elif not 0 <= hour <= 23:
        return None
    return {
        "minutes": hour * 60 + minute,
        "hour": hour,
        "minute": minute,
        "period": period,
        "ambiguous_12h": not period and 1 <= hour <= 12,
        "end": (colon or korean).end(),
    }


def _todo_deadline_end(clock, start_minutes):
    """12시간제 마감 표현을 현재 시각 뒤의 가장 가까운 시각으로 정한다."""
    if clock["ambiguous_12h"]:
        hour = clock["hour"] % 12
        minute = clock["minute"]
        candidates = [hour * 60 + minute, (hour + 12) * 60 + minute]
    else:
        candidates = [clock["minutes"]]
    future = [candidate for candidate in candidates if candidate > start_minutes]
    if future:
        return min(future)
    return min(candidates) + 24 * 60


def _schedule_from_todo(content, base_date, current_time=None):
    """투두 앞부분의 날짜·시간 표현을 일정 데이터로 변환한다."""
    text = (content or "").strip()
    if not text:
        return None

    base = datetime.strptime(base_date, "%Y-%m-%d")
    target_date = base_date
    date_match = re.match(
        r"^\s*(?:(\d{1,2})\s*월\s*(\d{1,2})\s*일|(\d{1,2})\s*[/\-]\s*(\d{1,2})|(\d{1,2})\s*일)(?=\s|$)",
        text)
    if date_match:
        month = int(date_match.group(1) or date_match.group(3) or base.month)
        day = int(date_match.group(2) or date_match.group(4) or date_match.group(5))
        try:
            target_date = datetime(base.year, month, day).strftime("%Y-%m-%d")
        except ValueError:
            return None
        text = text[date_match.end():].strip()

    now = current_time or datetime.now()
    start_time = end_time = ""
    end_date = target_date
    time_match = _todo_clock_match(text)
    if time_match:
        remainder = text[time_match["end"]:]
        range_separator = re.match(r"^\s*(?:부터|[~～-])\s*", remainder)
        deadline_suffix = re.match(r"^\s*까지", remainder)
        at_suffix = re.match(r"^\s*에(?=\s|$)", remainder)
        consumed = time_match["end"]
        if range_separator:
            second_offset = time_match["end"] + range_separator.end()
            end_match = _todo_clock_match(text[second_offset:])
            if not end_match:
                return None
            start_minutes = time_match["minutes"]
            end_minutes = end_match["minutes"]
            if time_match["period"] and not end_match["period"] and end_match["ambiguous_12h"]:
                end_minutes = (end_match["hour"] % 12 + (12 if time_match["period"] == "오후" else 0)) * 60 + end_match["minute"]
            if end_minutes <= start_minutes:
                end_minutes += 24 * 60
            start_time = f"{start_minutes // 60:02d}:{start_minutes % 60:02d}"
            end_time = f"{(end_minutes % (24 * 60)) // 60:02d}:{end_minutes % 60:02d}"
            end_date = (datetime.strptime(target_date, "%Y-%m-%d") + timedelta(days=end_minutes // (24 * 60))).strftime("%Y-%m-%d")
            consumed = second_offset + end_match["end"]
            suffix = re.match(r"^\s*까지", text[consumed:])
            if suffix:
                consumed += suffix.end()
        elif deadline_suffix:
            start_minutes = now.hour * 60 + now.minute
            end_minutes = _todo_deadline_end(time_match, start_minutes)
            start_time = f"{now.hour:02d}:{now.minute:02d}"
            end_time = f"{(end_minutes % (24 * 60)) // 60:02d}:{end_minutes % 60:02d}"
            end_date = (datetime.strptime(target_date, "%Y-%m-%d") + timedelta(days=end_minutes // (24 * 60))).strftime("%Y-%m-%d")
            consumed = time_match["end"] + deadline_suffix.end()
        else:
            start_minutes = time_match["minutes"]
            end_minutes = start_minutes + 60
            start_time = f"{start_minutes // 60:02d}:{start_minutes % 60:02d}"
            end_time = f"{(end_minutes % (24 * 60)) // 60:02d}:{end_minutes % 60:02d}"
            end_date = (datetime.strptime(target_date, "%Y-%m-%d") + timedelta(days=end_minutes // (24 * 60))).strftime("%Y-%m-%d")
            if at_suffix:
                consumed = time_match["end"] + at_suffix.end()
        text = text[consumed:].strip()

    if not date_match and not time_match:
        return None
    return {
        "title": text or content.strip(),
        "date": target_date,
        "start_date": target_date,
        "end_date": end_date,
        "start_time": start_time,
        "end_time": end_time,
    }


@app.route("/api/calendar/todos/delete", methods=["POST"])
@login_required
def api_calendar_todos_delete():
    tid = (request.json or {}).get("id")
    conn = get_db("calendar")
    conn.execute("DELETE FROM daily_todos WHERE id=? AND user=?", (tid, cur_user()))
    conn.commit()
    conn.close()
    return jsonify(ok=True)


@app.route("/api/calendar/week-plans")
@login_required
def api_calendar_week_plans():
    year = int(request.args.get("year") or datetime.now().year)
    month = int(request.args.get("month") or datetime.now().month)
    conn = get_db("calendar")
    rows = conn.execute("SELECT * FROM week_plans WHERE user=? AND year=? AND month=? ORDER BY week_no",
                        (cur_user(), year, month)).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))


@app.route("/api/calendar/week-plans/save", methods=["POST"])
@login_required
def api_calendar_week_plans_save():
    d = request.json or {}
    conn = get_db("calendar")
    conn.execute("""INSERT INTO week_plans(user,year,month,week_no,content,updated_at)
        VALUES(?,?,?,?,?,?) ON CONFLICT(user,year,month,week_no)
        DO UPDATE SET content=excluded.content,updated_at=excluded.updated_at""",
                 (cur_user(), int(d.get("year")), int(d.get("month")), int(d.get("week_no")),
                  d.get("content", ""), now_str()))
    conn.commit()
    conn.close()
    return jsonify(ok=True)

@app.route("/api/schedules/delete", methods=["POST"])
@login_required
def api_schedules_delete():
    sid = (request.json or {}).get("id")
    conn = get_db("calendar")
    old = conn.execute("SELECT * FROM schedules WHERE id=?", (sid,)).fetchone()
    conn.execute("DELETE FROM schedules WHERE id=?", (sid,))
    conn.commit(); conn.close()
    write_audit(cur_user(), "calendar", "delete", sid, before=dict(old) if old else None, detail="일정 삭제")
    return jsonify(ok=True)


# ════════════════════════════════════════════════════════════
#  프로젝트
# ════════════════════════════════════════════════════════════
@app.route("/api/projects")
@login_required
def api_projects():
    status = request.args.get("status", "open")
    conn = get_db("projects")
    rows = conn.execute("SELECT * FROM projects WHERE status=? ORDER BY created_at DESC", (status,)).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))

@app.route("/api/projects/save", methods=["POST"])
@login_required
def api_projects_save():
    d = request.json or {}
    emoji = d.get("emoji") or "📁"
    conn = get_db("projects")
    pid = d.get("id")
    if pid:
        conn.execute("UPDATE projects SET name=?,emoji=?,start_date=?,end_date=?,status=? WHERE id=?",
                     (d.get("name"), emoji, d.get("start_date"), d.get("end_date"), d.get("status", "open"), pid))
        conn.commit()
        write_audit(cur_user(), "projects", "update", pid, detail="프로젝트 수정")
    else:
        cur = conn.execute("INSERT INTO projects(name,emoji,start_date,end_date,status,created_at) VALUES(?,?,?,?,?,?)",
                           (d.get("name"), emoji, d.get("start_date"), d.get("end_date"), "open", now_str()))
        conn.commit(); pid = cur.lastrowid
        write_audit(cur_user(), "projects", "create", pid, after={"name": d.get("name")}, detail="프로젝트 등록")
    conn.close()
    return jsonify(ok=True, id=pid)


@app.route("/api/projects/status", methods=["POST"])
@login_required
def api_projects_status():
    d = request.json or {}
    status = d.get("status", "done")
    conn = get_db("projects")
    old = conn.execute("SELECT * FROM projects WHERE id=?", (d.get("id"),)).fetchone()
    conn.execute("UPDATE projects SET status=?,end_date=CASE WHEN ?='done' THEN COALESCE(NULLIF(end_date,''),?) ELSE end_date END WHERE id=?",
                 (status, status, datetime.now().strftime("%Y-%m-%d"), d.get("id")))
    conn.commit()
    conn.close()
    write_audit(cur_user(), "projects", "archive" if status == "done" else "restore",
                d.get("id"), before=dict(old) if old else None, after={"status": status},
                detail="프로젝트 완료 보관" if status == "done" else "프로젝트 복원")
    return jsonify(ok=True)


@app.route("/api/projects/delete", methods=["POST"])
@login_required
def api_projects_delete():
    pid = (request.json or {}).get("id")
    conn = get_db("projects")
    project = conn.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
    if not project:
        conn.close()
        return jsonify(ok=False, error="프로젝트를 찾을 수 없습니다."), 404
    card_ids = [r["id"] for r in conn.execute(
        "SELECT id FROM cards WHERE project_id=?", (pid,)
    ).fetchall()]
    if card_ids:
        marks = ",".join("?" for _ in card_ids)
        conn.execute(f"DELETE FROM comments WHERE card_id IN ({marks})", card_ids)
    conn.execute("DELETE FROM notifications WHERE project_id=?", (pid,))
    conn.execute("DELETE FROM cards WHERE project_id=?", (pid,))
    conn.execute("DELETE FROM projects WHERE id=?", (pid,))
    conn.commit()
    conn.close()
    write_audit(cur_user(), "projects", "delete", pid, before=dict(project),
                detail="프로젝트와 연결된 게시글·댓글 삭제")
    return jsonify(ok=True)

@app.route("/api/projects/<int:pid>/cards")
@login_required
def api_project_cards(pid):
    conn = get_db("projects")
    cards = conn.execute("SELECT * FROM cards WHERE project_id=? ORDER BY id", (pid,)).fetchall()
    cmts = conn.execute("""SELECT c.* FROM comments c JOIN cards k ON c.card_id=k.id
                           WHERE k.project_id=? ORDER BY c.id""", (pid,)).fetchall()
    conn.close()
    return jsonify(cards=rows_to_list(cards), comments=rows_to_list(cmts))

@app.route("/api/projects/card/add", methods=["POST"])
@login_required
def api_project_card_add():
    # 사진 등 첨부가 있으면 multipart, 없으면 JSON 모두 허용
    if request.content_type and "multipart" in request.content_type:
        pid = request.form.get("project_id")
        content = request.form.get("content", "")
        blocks_json = request.form.get("blocks_json") or "[]"
        imgs = []
        for f in request.files.getlist("images"):
            fn = save_upload(f, "card")
            if fn:
                imgs.append(fn)
        images = json.dumps(imgs, ensure_ascii=False) if imgs else None
    else:
        d = request.json or {}
        pid = d.get("project_id"); content = d.get("content"); images = None
        blocks_json = json.dumps(d.get("blocks") or [], ensure_ascii=False)
    conn = get_db("projects")
    cur = conn.execute("INSERT INTO cards(project_id,author,content,images,blocks_json,created_at) VALUES(?,?,?,?,?,?)",
                       (pid, cur_user(), content, images, blocks_json, now_str()))
    conn.commit(); cid = cur.lastrowid; conn.close()
    _create_mention_notifications(content, int(pid), "card", cid)
    write_audit(cur_user(), "projects", "add_card", cid, after={"content": content}, detail="프로젝트 게시글 추가")
    return jsonify(ok=True, id=cid)


@app.route("/api/projects/card/block/update", methods=["POST"])
@login_required
def api_project_card_block_update():
    d = request.json or {}
    conn = get_db("projects")
    card = conn.execute("SELECT * FROM cards WHERE id=?", (d.get("card_id"),)).fetchone()
    if not card:
        conn.close()
        return jsonify(ok=False, error="게시글을 찾을 수 없습니다."), 404
    try:
        blocks = json.loads(card["blocks_json"] or "[]")
        block = blocks[int(d.get("block_index"))]
        item = block["items"][int(d.get("item_index"))]
        if block.get("type") == "todo":
            item["done"] = bool(d.get("value"))
        elif block.get("type") == "progress":
            value = int(d.get("value") or 0)
            item["progress"] = value if value in (0, 25, 50, 75, 100) else 0
        else:
            raise ValueError("unsupported block")
    except (ValueError, TypeError, IndexError, KeyError, json.JSONDecodeError):
        conn.close()
        return jsonify(ok=False, error="진행 항목이 올바르지 않습니다."), 400
    conn.execute("UPDATE cards SET blocks_json=? WHERE id=?",
                 (json.dumps(blocks, ensure_ascii=False), d.get("card_id")))
    conn.commit()
    conn.close()
    write_audit(cur_user(), "projects", "update_progress", d.get("card_id"),
                after={"block_index": d.get("block_index"), "item_index": d.get("item_index"),
                       "value": d.get("value")}, detail="프로젝트 진행사항 변경")
    return jsonify(ok=True)

@app.route("/api/projects/comment/add", methods=["POST"])
@login_required
def api_project_comment_add():
    d = request.json or {}
    conn = get_db("projects")
    card = conn.execute("SELECT project_id FROM cards WHERE id=?", (d.get("card_id"),)).fetchone()
    cur = conn.execute("INSERT INTO comments(card_id,parent_id,author,content,created_at) VALUES(?,?,?,?,?)",
                       (d.get("card_id"), d.get("parent_id"), cur_user(), d.get("content"), now_str()))
    conn.commit(); cid = cur.lastrowid; conn.close()
    if card:
        _create_mention_notifications(d.get("content"), card["project_id"], "comment", cid)
    write_audit(cur_user(), "projects", "add_comment", cid, detail="프로젝트 댓글")
    return jsonify(ok=True, id=cid)


@app.route("/api/notifications")
@login_required
def api_notifications():
    conn = get_db("projects")
    rows = conn.execute("""SELECT n.*,p.name project_name,p.emoji
        FROM notifications n LEFT JOIN projects p ON p.id=n.project_id
        WHERE n.user=? AND n.dismissed=0 ORDER BY n.id DESC""", (cur_user(),)).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))


@app.route("/api/notifications/dismiss", methods=["POST"])
@login_required
def api_notifications_dismiss():
    d = request.json or {}
    conn = get_db("projects")
    if d.get("all"):
        conn.execute("UPDATE notifications SET dismissed=1 WHERE user=?", (cur_user(),))
    else:
        conn.execute("UPDATE notifications SET dismissed=1 WHERE id=? AND user=?",
                     (d.get("id"), cur_user()))
    conn.commit()
    conn.close()
    return jsonify(ok=True)


# ════════════════════════════════════════════════════════════
#  메모 (개인)
# ════════════════════════════════════════════════════════════
DEFAULT_MEMO_TAGS = [
    ("중요", "#e74c3c", 0),
    ("업무", "#2f80ed", 10),
    ("대기중", "#f2994a", 20),
    ("아이디어", "#9b51e0", 30),
    ("개인", "#27ae60", 40),
]


def _ensure_memo_tags(username):
    conn = get_db("memo")
    count = conn.execute(
        "SELECT COUNT(*) c FROM memo_tags WHERE author=?", (username,)
    ).fetchone()["c"]
    if not count:
        conn.executemany(
            "INSERT OR IGNORE INTO memo_tags(author,label,color,priority) VALUES(?,?,?,?)",
            [(username, label, color, priority) for label, color, priority in DEFAULT_MEMO_TAGS],
        )
        conn.commit()
    conn.close()


@app.route("/api/memo/tags")
@login_required
def api_memo_tags():
    _ensure_memo_tags(cur_user())
    conn = get_db("memo")
    rows = conn.execute(
        "SELECT label,color,priority FROM memo_tags WHERE author=? ORDER BY priority,label",
        (cur_user(),),
    ).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))


@app.route("/api/memo/tags/save", methods=["POST"])
@login_required
def api_memo_tags_save():
    tags = (request.json or {}).get("tags")
    if not isinstance(tags, list) or not tags or len(tags) > 20:
        return jsonify(ok=False, error="태그 목록을 확인해주세요."), 400
    cleaned = []
    labels = set()
    for index, item in enumerate(tags):
        if not isinstance(item, dict):
            return jsonify(ok=False, error="태그 형식이 올바르지 않습니다."), 400
        original = str(item.get("original") or item.get("label") or "").strip()[:20]
        label = str(item.get("label") or "").strip()[:20]
        color = str(item.get("color") or "").strip().lower()
        if not original or not label or label in labels:
            return jsonify(ok=False, error="태그 이름은 비어 있거나 중복될 수 없습니다."), 400
        if not re.match(r"^#[0-9a-f]{6}$", color):
            return jsonify(ok=False, error="태그 색상을 확인해주세요."), 400
        labels.add(label)
        cleaned.append((original, label, color, index * 10))
    conn = get_db("memo")
    try:
        # 이름 교환도 안전하도록 임시 이름을 거친다.
        for index, (original, _, _, _) in enumerate(cleaned):
            temp = f"__memo_tag_{index}_{int(time.time() * 1000)}__"
            conn.execute(
                "UPDATE memo_tags SET label=? WHERE author=? AND label=?",
                (temp, cur_user(), original),
            )
            conn.execute(
                "UPDATE memos SET tag=? WHERE author=? AND tag=?",
                (temp, cur_user(), original),
            )
            cleaned[index] = (temp, cleaned[index][1], cleaned[index][2], cleaned[index][3])
        for temp, label, color, priority in cleaned:
            conn.execute(
                "UPDATE memo_tags SET label=?,color=?,priority=? WHERE author=? AND label=?",
                (label, color, priority, cur_user(), temp),
            )
            conn.execute(
                "UPDATE memos SET tag=? WHERE author=? AND tag=?",
                (label, cur_user(), temp),
            )
        conn.commit()
    except sqlite3.IntegrityError:
        conn.rollback()
        conn.close()
        return jsonify(ok=False, error="태그 이름이 중복되었습니다."), 400
    conn.close()
    write_audit(cur_user(), "memo", "update_tags", detail="메모 태그 이름·색상·우선순위 변경")
    return jsonify(ok=True)


@app.route("/api/memo")
@login_required
def api_memo():
    status = request.args.get("status", "open")
    q = request.args.get("q", "")
    start = request.args.get("start", "")
    end = request.args.get("end", "")
    _ensure_memo_tags(cur_user())
    sort = request.args.get("sort", "manual")
    sql = """SELECT m.*,COALESCE(mt.color,'#64748b') tag_color,
             COALESCE(mt.priority,9999) tag_priority
             FROM memos m LEFT JOIN memo_tags mt
               ON mt.author=m.author AND mt.label=m.tag
             WHERE m.author=? AND m.status=?"""; p = [cur_user(), status]
    if q: sql += " AND content LIKE ?"; p.append(f"%{q}%")
    if start: sql += " AND created_at>=?"; p.append(start)
    if end: sql += " AND created_at<=?"; p.append(end + " 23:59:59")
    if status != "open":
        sql += " ORDER BY m.done_at DESC"
    elif sort == "manual":
        sql += " ORDER BY m.sort_order, m.id DESC"
    elif sort == "newest":
        sql += " ORDER BY m.id DESC"
    elif sort == "tag":
        sql += " ORDER BY m.tag, m.id DESC"
    else:
        sql += " ORDER BY tag_priority, m.sort_order, m.id DESC"
    conn = get_db("memo")
    rows = conn.execute(sql, p).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))

@app.route("/api/memo/add", methods=["POST"])
@login_required
def api_memo_add():
    d = request.json or {}
    conn = get_db("memo")
    cur = conn.execute("INSERT INTO memos(author,tag,content,status,sort_order,created_at) VALUES(?,?,?,?,?,?)",
                       (cur_user(), d.get("tag", ""), d.get("content"), "open", 0, now_str()))
    conn.commit(); mid = cur.lastrowid; conn.close()
    write_audit(cur_user(), "memo", "create", mid, after={"tag": d.get("tag"), "content": d.get("content")}, detail="메모 작성")
    return jsonify(ok=True, id=mid)


@app.route("/api/memo/update", methods=["POST"])
@login_required
def api_memo_update():
    d = request.json or {}
    conn = get_db("memo")
    old = conn.execute("SELECT * FROM memos WHERE id=? AND author=?",
                       (d.get("id"), cur_user())).fetchone()
    conn.execute("UPDATE memos SET tag=?,content=? WHERE id=? AND author=?",
                 (d.get("tag", ""), d.get("content", ""), d.get("id"), cur_user()))
    conn.commit()
    conn.close()
    write_audit(cur_user(), "memo", "update", d.get("id"),
                before=dict(old) if old else None,
                after={"tag": d.get("tag"), "content": d.get("content")},
                detail="메모 수정")
    return jsonify(ok=True)

@app.route("/api/memo/done", methods=["POST"])
@login_required
def api_memo_done():
    d = request.json or {}
    conn = get_db("memo")
    st = d.get("status", "done")
    conn.execute("UPDATE memos SET status=?, done_at=? WHERE id=? AND author=?",
                 (st, now_str() if st == "done" else None, d.get("id"), cur_user()))
    conn.commit(); conn.close()
    write_audit(cur_user(), "memo", "status", d.get("id"), after={"status": st}, detail="메모 상태 변경")
    return jsonify(ok=True)

@app.route("/api/memo/delete", methods=["POST"])
@login_required
def api_memo_delete():
    mid = (request.json or {}).get("id")
    conn = get_db("memo")
    old = conn.execute("SELECT * FROM memos WHERE id=? AND author=?", (mid, cur_user())).fetchone()
    conn.execute("DELETE FROM memos WHERE id=? AND author=?", (mid, cur_user()))
    conn.commit(); conn.close()
    write_audit(cur_user(), "memo", "delete", mid, before=dict(old) if old else None, detail="메모 삭제")
    return jsonify(ok=True)

@app.route("/api/memo/reorder", methods=["POST"])
@login_required
def api_memo_reorder():
    ids = (request.json or {}).get("ids", [])
    conn = get_db("memo")
    for i, mid in enumerate(ids):
        conn.execute("UPDATE memos SET sort_order=? WHERE id=? AND author=?", (i, mid, cur_user()))
    conn.commit(); conn.close()
    return jsonify(ok=True)


# ════════════════════════════════════════════════════════════
#  공지
# ════════════════════════════════════════════════════════════
@app.route("/api/notices")
@login_required
def api_notices():
    conn = get_db("notice")
    rows = conn.execute("SELECT * FROM notices ORDER BY id DESC").fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))

@app.route("/api/notices/add", methods=["POST"])
@admin_required
def api_notices_add():
    d = request.json or {}
    conn = get_db("notice")
    cur = conn.execute(
        "INSERT INTO notices(author,title,content,level,created_at,status,completed_at) VALUES(?,?,?,?,?,?,?)",
        (cur_user(), d.get("title"), d.get("content"), d.get("level", "normal"),
         now_str(), "open", None))
    conn.commit(); nid = cur.lastrowid; conn.close()
    write_audit(cur_user(), "notice", "create", nid, after={"title": d.get("title")}, detail="공지 작성")
    return jsonify(ok=True, id=nid)


@app.route("/api/notices/status", methods=["POST"])
@admin_required
def api_notices_status():
    d = request.json or {}
    nid = d.get("id")
    status = "done" if d.get("status") == "done" else "open"
    completed_at = now_str() if status == "done" else None
    conn = get_db("notice")
    conn.execute("UPDATE notices SET status=?,completed_at=? WHERE id=?",
                 (status, completed_at, nid))
    conn.commit()
    conn.close()
    write_audit(cur_user(), "notice", "status", nid, after={"status": status},
                detail="공지 상태 변경")
    return jsonify(ok=True)

@app.route("/api/notices/delete", methods=["POST"])
@admin_required
def api_notices_delete():
    nid = (request.json or {}).get("id")
    conn = get_db("notice")
    conn.execute("DELETE FROM notices WHERE id=?", (nid,))
    conn.commit(); conn.close()
    write_audit(cur_user(), "notice", "delete", nid, detail="공지 삭제")
    return jsonify(ok=True)


# ════════════════════════════════════════════════════════════
#  통합 검색 (F4)
# ════════════════════════════════════════════════════════════
@app.route("/api/search")
@login_required
def api_search():
    q = request.args.get("q", "").strip()
    out = {"contacts": [], "items": [], "tx": [], "schedules": [], "memo": []}
    if not q:
        return jsonify(out)
    like = f"%{q}%"
    c = get_db("contacts"); out["contacts"] = rows_to_list(c.execute(
        "SELECT id,name,biz_no FROM contacts WHERE name LIKE ? OR biz_no LIKE ? LIMIT 10", (like, like)).fetchall()); c.close()
    it = get_db("items"); out["items"] = rows_to_list(it.execute(
        "SELECT id,sku,name FROM items WHERE sku LIKE ? OR name LIKE ? LIMIT 10", (like, like)).fetchall()); it.close()
    tx = get_db("transactions"); out["tx"] = rows_to_list(tx.execute(
        "SELECT id,date,type,item_name,total FROM transactions WHERE item_name LIKE ? OR contact LIKE ? OR sku LIKE ? LIMIT 10",
        (like, like, like)).fetchall()); tx.close()
    ca = get_db("calendar"); out["schedules"] = rows_to_list(ca.execute(
        "SELECT id,title,start_date FROM schedules WHERE title LIKE ? OR memo LIKE ? LIMIT 10", (like, like)).fetchall()); ca.close()
    me = get_db("memo"); out["memo"] = rows_to_list(me.execute(
        "SELECT id,tag,content FROM memos WHERE author=? AND content LIKE ? LIMIT 10", (cur_user(), like)).fetchall()); me.close()
    return jsonify(out)


@app.route("/api/search/all")
@login_required
def api_search_all():
    q = request.args.get("q", "").strip()
    if not q:
        return jsonify(total=0, categories=[])
    like = f"%{q}%"
    categories = []

    def add_category(name, icon, items):
        if items:
            categories.append({"name": name, "icon": icon, "items": items})

    conn = get_db("contacts")
    rows = conn.execute(
        """SELECT id,name,biz_no,ceo,phone,address FROM contacts
           WHERE name LIKE ? OR biz_no LIKE ? OR ceo LIKE ? OR phone LIKE ?
              OR address LIKE ? OR memo LIKE ? ORDER BY name LIMIT 100""",
        (like,) * 6,
    ).fetchall()
    conn.close()
    add_category("거래처", "🏬", [{
        "title": r["name"], "meta": " · ".join(x for x in (r["biz_no"], r["ceo"], r["phone"]) if x),
        "target": {"tab": "contacts", "kind": "contact", "id": r["id"]},
    } for r in rows])

    conn = get_db("items")
    rows = conn.execute(
        """SELECT id,sku,name,spec,unit,location FROM items
           WHERE sku LIKE ? OR name LIKE ? OR spec LIKE ? OR unit LIKE ? OR location LIKE ?
              OR distributor LIKE ? OR manufacturer LIKE ? ORDER BY sku LIMIT 100""",
        (like,) * 7,
    ).fetchall()
    conn.close()
    add_category("품목", "📦", [{
        "title": f"{r['sku']} · {r['name']}", "meta": " · ".join(x for x in (r["spec"], r["unit"], r["location"]) if x),
        "target": {"tab": "items", "kind": "item", "id": r["id"]},
    } for r in rows])

    conn = get_db("transactions")
    rows = conn.execute(
        """SELECT id,date,type,contact,sku,item_name,total,source_voucher_id FROM transactions
           WHERE date LIKE ? OR type LIKE ? OR contact LIKE ? OR sku LIKE ? OR item_name LIKE ?
              OR spec LIKE ? OR note LIKE ? OR channel LIKE ? OR order_no LIKE ?
              OR shipment_no LIKE ? OR project_ref LIKE ?
           ORDER BY date DESC,id DESC LIMIT 100""",
        (like,) * 11,
    ).fetchall()
    conn.close()
    add_category("매입·매출", "🧾", [{
        "title": f"{r['date']} [{r['type']}] {r['item_name'] or r['sku'] or ''}",
        "meta": f"{r['contact'] or '-'} · {r['total'] or 0:,.0f}원",
        "target": {"tab": "inout", "kind": "tx", "id": r["id"]},
    } for r in rows])

    conn = get_db("vouchers")
    rows = conn.execute(
        """SELECT v.id,v.kind,v.date,v.seq,v.contact,v.total,v.memo,
                  GROUP_CONCAT(COALESCE(i.name,'') || ' ' || COALESCE(i.spec,''),' · ') item_summary
           FROM vouchers v LEFT JOIN voucher_items i ON i.voucher_id=v.id
           GROUP BY v.id
           HAVING v.date LIKE ? OR v.contact LIKE ? OR v.memo LIKE ?
              OR COALESCE(item_summary,'') LIKE ?
           ORDER BY v.date DESC,v.id DESC LIMIT 100""",
        (like, like, like, like),
    ).fetchall()
    conn.close()
    kind_names = {"sale": "판매", "purchase": "구매", "quote": "견적"}
    add_category("판매·견적", "📄", [{
        "title": f"{r['date']} -{r['seq']} {r['contact'] or ''}",
        "meta": f"{kind_names.get(r['kind'], r['kind'])} · {r['item_summary'] or ''} · {r['total'] or 0:,.0f}원",
        "target": {"tab": "voucher", "kind": "voucher", "id": r["id"], "voucher_kind": r["kind"]},
    } for r in rows if r["kind"] != "purchase"])

    conn = get_db("calendar")
    rows = conn.execute(
        """SELECT id,title,start_date,end_date,memo FROM schedules
           WHERE title LIKE ? OR memo LIKE ? OR start_date LIKE ? OR end_date LIKE ?
           ORDER BY start_date DESC,id DESC LIMIT 100""",
        (like, like, like, like),
    ).fetchall()
    conn.close()
    add_category("일정", "📅", [{
        "title": r["title"], "meta": f"{r['start_date'] or ''}{' ~ ' + r['end_date'] if r['end_date'] and r['end_date'] != r['start_date'] else ''}",
        "target": {"tab": "calendar", "kind": "schedule", "id": r["id"]},
    } for r in rows])

    conn = get_db("projects")
    rows = conn.execute(
        """SELECT p.id,p.name,p.emoji,p.status,MAX(c.created_at) updated_at
           FROM projects p LEFT JOIN cards c ON c.project_id=p.id
           WHERE p.name LIKE ? OR c.content LIKE ?
           GROUP BY p.id ORDER BY COALESCE(updated_at,p.created_at) DESC LIMIT 100""",
        (like, like),
    ).fetchall()
    conn.close()
    add_category("프로젝트", "📁", [{
        "title": f"{r['emoji'] or '📁'} {r['name']}", "meta": "완료" if r["status"] == "done" else "진행중",
        "target": {"tab": "projects", "kind": "project", "id": r["id"], "status": r["status"]},
    } for r in rows])

    conn = get_db("memo")
    rows = conn.execute(
        """SELECT id,tag,content,status,created_at FROM memos
           WHERE author=? AND (tag LIKE ? OR content LIKE ?)
           ORDER BY created_at DESC LIMIT 100""",
        (cur_user(), like, like),
    ).fetchall()
    conn.close()
    add_category("개인 메모", "📝", [{
        "title": (r["content"] or "")[:100], "meta": f"{r['tag'] or '태그 없음'} · {r['status']}",
        "target": {"tab": "memo", "kind": "memo", "id": r["id"], "status": r["status"]},
    } for r in rows])

    conn = get_db("notice")
    rows = conn.execute(
        """SELECT id,title,content,level,created_at FROM notices
           WHERE title LIKE ? OR content LIKE ? OR author LIKE ?
           ORDER BY created_at DESC LIMIT 100""",
        (like, like, like),
    ).fetchall()
    conn.close()
    add_category("공지사항", "📢", [{
        "title": r["title"], "meta": (r["content"] or "")[:120],
        "target": {"tab": "notices", "kind": "notice", "id": r["id"]},
    } for r in rows])

    settings_index = [
        ("회사 설정", "회사명 사업자등록번호 대표자명 주소 전화번호 팩스 권한 DB", "company", "set-company"),
        ("출력 디자인", "견적서 명세표 로고 워터마크 도장 홈페이지 계좌번호 인쇄", "print", "set-print-color"),
        ("내 계정", "아이디 비밀번호 이름 이메일 찾기 질문", "account", "ac-user"),
        ("단축키", "F1 F2 F3 F4 F6 F8 F9 F10 사이드바", "shortcuts", ""),
        ("빠른실행", "상단 입고 출고 검색 메모 프로젝트 거래처", "quick", ""),
        ("멀티 검색", "일반검색 전체검색 기본 시작 탭 검색 사이트 홈택스 바로가기 링크", "search", "multi-search-settings-panel"),
        ("테이블 관리", "헤더 설정 열 표시 숨김 거래처 품목 매입매출 재고 입출고 전표", "tables", "tbl-mgmt-body"),
    ]
    q_lower = q.lower()
    setting_results = [{
        "title": title, "meta": keywords,
        "target": {"tab": "settings", "kind": "settings", "settings_cat": cat, "element": element},
    } for title, keywords, cat, element in settings_index if q_lower in f"{title} {keywords}".lower()]
    add_category("설정", "⚙️", setting_results)

    return jsonify(total=sum(len(x["items"]) for x in categories), categories=categories)


if __name__ == "__main__":
    bootstrap()
    print("재고 관리 프로그램  ->  http://localhost:5000  (사내 접속: http://<서버IP>:5000)")
    app.run(host="0.0.0.0", port=5000, debug=False)
